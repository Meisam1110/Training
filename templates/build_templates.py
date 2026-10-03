#!/usr/bin/env python3
"""Regenerates templates/*.docx from the wording in backend/app/documents/wording.py.

    python templates/build_templates.py

The templates carry only ``{{placeholders}}`` (no real names, IDs, phones, e-mails, numbers), no macros
and no mail-merge data-source links. Images are the reusable Irancell-look assets in templates/assets/.
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "backend"))

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402,F401
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from app.documents.specs import SPECS  # noqa: E402
from app.documents.wording import DOCS  # noqa: E402

ASSETS = HERE / "assets"
FONT = "Tahoma"
YELLOW, BLUE, GREY = "FFCB05", RGBColor(0x00, 0x4A, 0x8F), "BFBFBF"
CONTENT_W_CM = 17.0


# --------------------------------------------------------------------------- low-level helpers
def _rfonts(rpr, font=FONT):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)


def setup(doc: Document, header_img: str | None = None, footer_img: str | None = None) -> None:
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    _rfonts(rpr)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US"); lang.set(qn("w:bidi"), "fa-IR")
    rpr.append(lang)
    szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), "21"); rpr.append(szcs)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    if header_img:
        sec.header_distance = Cm(0.5)
        sec.top_margin = Cm(4.7)
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.add_run().add_picture(str(ASSETS / header_img), width=Cm(19.0))
    if footer_img:
        sec.footer_distance = Cm(0.5)
        sec.bottom_margin = Cm(2.4)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run().add_picture(str(ASSETS / footer_img), width=Cm(19.0))


def _mark_rtl_run(run):
    rpr = run._r.get_or_add_rPr()
    _rfonts(rpr)
    rpr.append(OxmlElement("w:rtl"))
    if run.bold:
        rpr.append(OxmlElement("w:bCs"))


def para(container, text: str, *, rtl: bool, bold=False, size=None, color=None, center=False,
         space_after=6, italic=False, keep_font=True):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(space_after), Pt(0)
    if rtl:
        ppr = p._p.get_or_add_pPr()
        ppr.append(OxmlElement("w:bidi"))
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold or None, italic or None
        if size:
            r.font.size = Pt(size)
            szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2)))
            r._r.get_or_add_rPr().append(szcs)
        if color is not None:
            r.font.color.rgb = color
        if rtl:
            _mark_rtl_run(r)
    return p


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), GREY)
        b.append(e)
    tblpr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "60"); e.set(qn("w:type"), "dxa"); mar.append(e)
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "100"); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblpr.append(mar)


def table(doc, ncols: int, rtl: bool, widths_cm: list[float]):
    t = doc.add_table(rows=0, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _borders(t)
    if rtl:
        t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    t._widths = widths_cm
    for col, w in zip(t.columns, widths_cm):
        col.width = Cm(w)
    return t


def add_row(t, texts: list[str], rtl: bool, *, bold=False, fill=None, header=False):
    row = t.add_row()
    for i, (cell, text) in enumerate(zip(row.cells, texts)):
        cell.width = Cm(t._widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if rtl:
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        r = p.add_run(text)
        r.bold = bold or None
        r.font.size = Pt(9.5)
        if rtl:
            _mark_rtl_run(r)
        if fill:
            _shade(cell, fill)
    if header:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    return row


def spacer(doc, pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    return p


# --------------------------------------------------------------------------- block renderer
def emit(doc, blocks, rtl: bool):
    for b in blocks:
        kind = b[0]
        if kind == "banner":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ASSETS / b[1]), width=Cm(CONTENT_W_CM))
        elif kind == "image":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ASSETS / b[1]), width=Cm(b[2]))
        elif kind == "title":
            para(doc, b[1], rtl=rtl, bold=True, size=15, color=BLUE, space_after=10)
        elif kind == "p":
            para(doc, b[1], rtl=rtl)
        elif kind == "note":
            para(doc, b[1], rtl=rtl, size=9, italic=True)
        elif kind == "link":
            label, url = b[1], b[2]
            para(doc, url if label == url else f"{label}: {url}", rtl=rtl, color=RGBColor(0x05, 0x63, 0xC1))
        elif kind == "bullets":
            for t in b[1]:
                para(doc, "•  " + t if not rtl else "•  " + t, rtl=rtl, space_after=3)
            spacer(doc, 2)
        elif kind == "meta":
            for label, val in b[1]:
                para(doc, f"{label}: {val}", rtl=rtl, space_after=2)
            spacer(doc, 4)
        elif kind == "kv":
            t = table(doc, 2, rtl, [4.5, CONTENT_W_CM - 4.5])
            for label, val in b[1]:
                row = add_row(t, [label, val], rtl, fill=None)
                _shade(row.cells[0], YELLOW)
                for r in row.cells[0].paragraphs[0].runs:
                    r.bold = True
            spacer(doc)
        elif kind == "grid":
            cols = b[2]
            n = len(cols)
            if n == 5 and cols[0][1].endswith("._n}}"):
                widths = [1.2, 2.6, 6.6, 3.6, 3.0]
            elif n == 4 and cols[0][1].endswith("._n}}"):
                widths = [1.3, 3.0, CONTENT_W_CM - 1.3 - 3.0 - 4.2, 4.2]
            elif n == 4:
                widths = [CONTENT_W_CM * f for f in (0.22, 0.17, 0.18, 0.43)]
            else:
                widths = [CONTENT_W_CM / n] * n
            t = table(doc, n, rtl, widths)
            add_row(t, [h for h, _ in cols], rtl, bold=True, fill=YELLOW, header=True)
            add_row(t, [c for _, c in cols], rtl)
            spacer(doc)
        elif kind == "sign":
            for i, ln in enumerate(b[1]):
                sp = para(doc, ln, rtl=rtl, bold=("{{signer_name_fa}}" in ln or "{{unit_fa}}" in ln), space_after=1)
                sp.paragraph_format.keep_with_next = i < len(b[1]) - 1      # a signature block never splits across pages
                sp.paragraph_format.keep_together = True
        elif kind == "sign_formal":
            # official letter: closing + signature block on the left side of the page (right cell empty in RTL)
            t = table(doc, 2, rtl, [CONTENT_W_CM - 7.0, 7.0])
            t._tbl.tblPr.remove(t._tbl.tblPr.find(qn("w:tblBorders")))
            row = t.add_row()
            cell = row.cells[1]
            row.cells[0].width, cell.width = Cm(CONTENT_W_CM - 7.0), Cm(7.0)
            cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
            for ln in b[1]:
                para(cell, ln, rtl=rtl, center=True, bold=("{{signer_name_fa}}" in ln or "{{unit_fa}}" in ln), space_after=1)
            if not cell.paragraphs:
                cell.add_paragraph()
        else:
            raise ValueError(kind)


def add_lang_marker(doc, lang: str):
    p = doc.add_paragraph("[[LANG:%s]]" % lang)
    p.paragraph_format.space_after = Pt(0)


def build(kind: str, out: Path):
    spec = SPECS[kind]
    letter = kind == "institute_letter"
    doc = Document()
    setup(doc, "letterhead_top.jpg" if letter else None, "letterhead_bottom.jpg" if letter else None)
    sections = DOCS[kind]
    if spec.bilingual:
        for i, lang in enumerate(("fa", "en")):
            add_lang_marker(doc, lang)
            if i:
                p = doc.add_paragraph()
                p.add_run().add_break(WD_BREAK.PAGE)
            emit(doc, sections[lang], rtl=(lang == "fa"))
    else:
        (lang, blocks), = sections.items()
        emit(doc, blocks, rtl=(lang == "fa"))
    cp = doc.core_properties
    cp.author = cp.last_modified_by = "TMS"
    cp.created = cp.modified = __import__("datetime").datetime(2025, 1, 1)
    cp.revision, cp.category, cp.keywords, cp.subject = 1, "", "", ""
    cp.author = "TMS"
    doc.core_properties.title = spec.title
    doc.core_properties.comments = "Template with {{placeholders}}; generated by templates/build_templates.py"
    doc.save(out)


def main():
    for kind, spec in SPECS.items():
        out = HERE / spec.template
        build(kind, out)
        print("built", out.name)


if __name__ == "__main__":
    main()
