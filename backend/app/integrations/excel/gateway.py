"""ExcelGateway: reads the iCare IDP export, the IFS employee sheet and the legacy workbook
(sheets IFS, IDP, Record, CHistory); writes the legacy layout back.

Design rules
* The source file is read once into memory and closed immediately (no lock is held while parsing,
  the file is never written to). ``openpyxl`` runs in ``read_only`` streaming mode.
* Error messages carry row numbers and column names only - never cell values (PII).
* Everything is a pure function of the file content, so re-reading is idempotent.
"""
from __future__ import annotations

import io
import logging
import time
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from functools import partial
from pathlib import Path
from typing import Any, Callable, Iterator, Sequence

from openpyxl import Workbook, load_workbook

from app.contracts import EmployeeRecord, IdpRow, LegacyData, RecordRow, RunRecord, SyncReport

from . import normalize as N

log = logging.getLogger(__name__)

MAX_HEADER_SCAN = 40

IDP_ALIASES = {
    "emp_no": ["Employee ID", "Emp No", "Employee Number"],
    "reporter_name": ["Reporter: Name", "Reporter Name", "Reporter"],
    "job_level": ["Job Level"],
    "department": ["Reporter Department", "Department"],
    "summary": ["Summary", "Course Name"],
    "delivery_method": ["Delivery Method"],
    "current_status": ["Current Status", "Status"],
    "idp_type": ["IDP Type"],
    "course_category": ["Course Category", "Category"],
    "attendance_status": ["Attendance Status"],
    "start_date": ["Course Start Date", "Start Date"],
    "end_date": ["Course End Date", "End Date"],
    "ticket": ["Ticket", "Key", "Issue key"],
    "section_b": ["Section B", "SectionB"],
}
IDP_REQUIRED = ("emp_no", "summary", "ticket")

IFS_ALIASES = {
    "emp_no": ["Emp No", "Employee ID", "Employee Number"],
    "status": ["Employee Status", "Status"],
    "org_code": ["Org Code"],
    "org_name": ["Organization Name"],
    "cost_center": ["Cost Center"],
    "first_en": ["Fname", "First Name"],
    "last_en": ["Lname", "Last Name"],
    "gender": ["Gender"],
    "mobile": ["Mobile"],
    "email": ["Email"],
    "first_fa": ["Farsi First Name"],
    "last_fa": ["Farsi Last Name"],
    "position": ["Position Title"],
    "sup_emp_no": ["Sup Emp No"],
    "national_code": ["National Code"],
    "division": ["Division"],
}
IFS_REQUIRED = ("emp_no", "first_en", "last_en")

RECORD_ALIASES = {
    "emp_no": ["Employee Id", "ID"],
    "ticket": ["Ticket"],
    "course_name": ["Course Name"],
    "institute": ["Institute"],
    "start_date": ["start Date"],
    "motivation": ["Motivation"],
    "registration": ["Registration"],
    "calendar": ["Calendar"],
    "payment": ["Payment"],
    "comment": ["Comment"],
}
RECORD_REQUIRED = ("emp_no", "ticket")

RUN_ALIASES = {
    "course_name": ["Course Name"],
    "course_name_fa": ["Course NameFarsi", "Course Name Farsi"],
    "type": ["Type"],
    "institute": ["Institute"],
    "duration": ["Duration"],
    "start_date": ["start Date"],
    "end_date": ["End Date"],
    "days": ["Day"],
    "time": ["Time"],
    "sessions": ["Sessions Count"],
    "send_date": ["Send Date"],
    "deadline": ["Deadline"],
    "amount": ["Amount"],
    "section_b": ["SectionB", "Section B"],
    "link": ["Link"],
}
RUN_REQUIRED = ("course_name",)

# Legacy output layout (column order as in the workbook).
LEGACY_IFS_COLS = ["Emp No", "Employee Status", "Employment Type", "Org Code", "Organization Name", "Cost Center",
                   "Fname", "Lname", "Gender", "Mobile", "Email", "Farsi First Name", "Farsi Last Name",
                   "Position Title", "Job Level", "Company Office", "Sup Emp No", "Sup Emp Name", "Sup Emp Mobile",
                   "Sup Emp Emaile", "Pic", "Education", "Date Of Birth", "Start Date", "Position Code", "Job Id",
                   "Psychological Char", "Value String", "Id No", "National Code", "Division"]
LEGACY_IDP_COLS = ["Employee ID", "Reporter:  Name", "Job Level", "Reporter Department", "Summary", "Delivery Method",
                   "Current Status", "IDP Type", "Course Category", "Attendance Status", "Course Start Date",
                   "Course End Date", "Ticket", "Section B"]
LEGACY_RECORD_COLS = ["Row", "Employee Id", "Fullname", "division", "Email", "sup Email", "Ticket",
                      "Course In same tim", "Course Name", "Course Name2", "Type", "Institute", "Duration",
                      "jStart Date", "End Date", "start Date", "Time", "Sessions Count", "Day", "Send Date",
                      "Deadline", "insutute FA Name", "Motivation", "SectionB", "Registration", "Calendar",
                      "Payment", "Course Time line", "Days toEnd date", "helper", "helper2", "helper3", "Helper4",
                      "Key", "Comment", "link", "Amount"]
LEGACY_CHISTORY_COLS = ["Row", "Course Name", "Course NameFarsi", "Type", "Institute", "Duration", "jStart Date",
                        "End Date", "start Date", "Time", "Sessions Count", "Day", "Send Date", "Deadline",
                        "Email Sent to", "No of Motivations", "Percent", "SectionB", "Registration", "Calendar",
                        "Payment Process", "Invoice Receive Date", "PR process Start Date", "Amount", "Certificate",
                        "Link", "Days To course", "End date2", "Key"]


@dataclass(frozen=True)
class ExcelLimits:
    """Resource guards against hostile / accidental huge workbooks."""
    max_file_bytes: int = 100 * 1024 * 1024          # compressed size read into memory
    max_uncompressed_bytes: int = 250 * 1024 * 1024  # sum of all zip members after decompression
    max_members: int = 2000
    max_ratio: float = 200.0                          # uncompressed / compressed, checked above 20 MB
    max_rows: int = 200_000                           # rows scanned per sheet
    max_cols: int = 200
    max_cell_chars: int = 2000
    max_seconds: float | None = None                  # wall-clock budget per sheet read


class ExcelFormatError(ValueError):
    """The workbook does not look like the expected export (message never contains cell values)."""


class ExcelReadAborted(ExcelFormatError):
    """Read stopped by the caller (cancel callback) or by the time budget."""


class _Guard:
    def __init__(self, limits: ExcelLimits, should_cancel: Callable[[], bool] | None):
        self.limits, self.should_cancel = limits, should_cancel
        self.started = time.monotonic()
        self.row_cap_hit = False

    def tick(self, n: int) -> None:
        if n % 500 == 0:
            if self.should_cancel is not None and self.should_cancel():
                raise ExcelReadAborted("read cancelled by the caller")
            if self.limits.max_seconds is not None and time.monotonic() - self.started > self.limits.max_seconds:
                raise ExcelReadAborted("read exceeded its time budget")


def _read_bytes(path: str | Path, limits: ExcelLimits) -> bytes:
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"Excel file not found: {p.name}")
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelFormatError(f"Unsupported file type '{p.suffix}' (use .xlsx or .xlsm)")
    if p.stat().st_size > limits.max_file_bytes:
        raise ExcelFormatError("Workbook file is too large")
    with open(p, "rb") as f:          # read-only, closed immediately: no lock kept on the source
        return f.read()


def check_zip_bomb(data: bytes, limits: ExcelLimits) -> None:
    """Reject before parsing: inspects only the zip central directory (cheap, nothing is decompressed)."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except zipfile.BadZipFile:
        raise ExcelFormatError("Cannot open workbook (not a valid .xlsx/.xlsm archive)") from None
    if len(infos) > limits.max_members:
        raise ExcelFormatError("Workbook has too many internal parts")
    total = sum(i.file_size for i in infos)
    if total > limits.max_uncompressed_bytes:
        raise ExcelFormatError("Workbook is too large when decompressed (possible zip bomb)")
    if total > 20 * 1024 * 1024 and total / max(1, len(data)) > limits.max_ratio:
        raise ExcelFormatError("Workbook compression ratio is implausibly high (possible zip bomb)")
    if any(i.filename.startswith("/") or ".." in i.filename.split("/") for i in infos):
        raise ExcelFormatError("Workbook contains unsafe member paths")


def _open(path: str | Path, limits: ExcelLimits):
    data = _read_bytes(path, limits)
    check_zip_bomb(data, limits)
    try:
        return load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:  # noqa: BLE001 - zip / xml corruption
        raise ExcelFormatError(f"Cannot open workbook ({type(exc).__name__})") from None


def _rows(ws, guard: _Guard | None = None) -> Iterator[tuple[int, Sequence[Any]]]:
    """Yield (1-based row number, values) streaming; repairs bogus dimension metadata."""
    try:
        ws.reset_dimensions()
    except Exception:  # noqa: BLE001
        pass
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if guard is not None:
            guard.tick(i)
            if i > guard.limits.max_rows:
                guard.row_cap_hit = True
                return
            if len(row) > guard.limits.max_cols:
                row = row[: guard.limits.max_cols]
        yield i, row


def _find_sheet(wb, name: str | None):
    if name:
        for ws in wb.worksheets:
            if ws.title.strip().lower() == name.lower():
                return ws
        return None
    return wb.worksheets[0] if wb.worksheets else None


class _Table:
    """A located header + the iterator of data rows after it."""

    def __init__(self, ws, aliases, required, sheet_label: str, min_score: int = 3, guard: _Guard | None = None):
        self.cmap: dict[str, int] = {}
        self.header_row = 0
        self.label = sheet_label
        self._ws = ws
        self._aliases = aliases
        self._required = required
        self._min = min_score
        self._buffer: list[tuple[int, Sequence[Any]]] = []
        self.guard = guard
        self.notes: list[str] = []
        self.header_names: list[str] = []
        self._it = _rows(ws, guard)
        self._locate()

    def _locate(self) -> None:
        best_score, best_row, best_map = 0, 0, {}
        seen: list[tuple[int, Sequence[Any]]] = []
        for n, row in self._it:
            seen.append((n, row))
            sc = N.score_header(row, self._aliases)
            if sc > best_score:
                best_score, best_row, best_map = sc, n, N.build_column_map(row, self._aliases)
            if all(k in best_map for k in self._required) and sc >= len(self._aliases) - 3 or n >= MAX_HEADER_SCAN:
                break
        missing = [self._aliases[k][0] for k in self._required if k not in best_map]
        if best_score < self._min or missing:
            raise ExcelFormatError(
                f"Sheet '{self.label}': header row not found; missing columns: {', '.join(missing) or 'many'}")
        self.cmap, self.header_row = best_map, best_row
        self.header_names = [N.to_text(c) for c in next((r for n, r in seen if n == best_row), [])]
        # data starts after the header; duplicate header rows (legacy has two) are skipped later
        self._buffer = [(n, r) for n, r in seen if n > best_row]

    def data(self) -> Iterator[tuple[int, Sequence[Any]]]:
        key_col = next(iter(self._cmap_first()))
        header_text = N.norm_header(self._aliases[key_col][0])
        aliases_norm = {N.norm_header(a) for a in self._aliases[key_col]}
        for n, row in self._chain():
            first = N.norm_header(N.cell(row, self.cmap, key_col))
            if first in aliases_norm or first == header_text:
                continue  # repeated header row
            yield n, self._clip(n, row)
        if self.guard is not None and self.guard.row_cap_hit:
            self.notes.append(f"{self.label}: row limit of {self.guard.limits.max_rows} reached; remaining rows ignored")

    def _clip(self, n: int, row: Sequence[Any]) -> Sequence[Any]:
        cap = self.guard.limits.max_cell_chars if self.guard is not None else 0
        if not cap:
            return row
        out = None
        for i, v in enumerate(row):
            if isinstance(v, str) and len(v) > cap:
                out = list(row) if out is None else out
                out[i] = v[:cap]
                col = self.header_names[i] if i < len(self.header_names) and self.header_names[i] else f"#{i + 1}"
                self.notes.append(f"{self.label} row {n}: value in column '{col}' was over-long and truncated")
        return out if out is not None else row

    def _cmap_first(self):
        return [k for k in self._required if k in self.cmap] or list(self.cmap)

    def _chain(self):
        yield from self._buffer
        yield from self._it


class OpenpyxlExcelGateway:
    """Implements ``contracts.ExcelGateway`` plus ``*_report`` variants returning a ``SyncReport``.

    Safety: read-only streaming (macros in .xlsm are never loaded or executed: ``keep_vba`` is off),
    formula cells are read as their cached VALUES (never evaluated; never-calculated formulas read as
    empty), zip-bomb / size / row / cell caps from ``ExcelLimits``, optional ``should_cancel`` hook and
    time budget, and a per-row try/except so one bad row is reported (row number only) instead of aborting.
    """

    def __init__(self, limits: ExcelLimits | None = None, should_cancel: Callable[[], bool] | None = None):
        self.limits = limits or ExcelLimits()
        self.should_cancel = should_cancel

    def _guard(self) -> _Guard:
        return _Guard(self.limits, self.should_cancel)

    # ---------------------------------------------------------------- IDP
    def read_idp(self, path: str) -> list[IdpRow]:
        return self.read_idp_report(path)[0]

    def read_idp_report(self, path: str, sheet: str | None = None) -> tuple[list[IdpRow], SyncReport]:
        wb = _open(path, self.limits)
        try:
            ws = _find_sheet(wb, sheet) if sheet else self._pick(wb, "idp", IDP_ALIASES, IDP_REQUIRED)
            if ws is None:
                raise ExcelFormatError("IDP sheet not found")
            return self._parse_idp(ws)
        finally:
            wb.close()

    def _pick(self, wb, preferred: str, aliases, required):
        """Prefer the sheet named *preferred*; otherwise the first sheet whose header matches."""
        cand = _find_sheet(wb, preferred)
        if cand is not None:
            return cand
        for ws in wb.worksheets:
            try:
                _Table(ws, aliases, required, ws.title, guard=self._guard())
                return ws
            except ExcelFormatError:
                continue
        return wb.worksheets[0] if wb.worksheets else None

    def _parse_idp(self, ws) -> tuple[list[IdpRow], SyncReport]:
        report = SyncReport(source="excel:idp", started_at=datetime.now(timezone.utc))
        tab = _Table(ws, IDP_ALIASES, IDP_REQUIRED, ws.title, guard=self._guard())
        out: list[IdpRow] = []
        seen: set[tuple[str, str]] = set()
        for n, row in tab.data():
            try:
                g = partial(N.cell, row, tab.cmap)
                emp = N.clean_emp_no(g("emp_no"))
                ticket = N.clean_ticket(g("ticket"))
                if not emp and not ticket:
                    if any(not N.is_blank(c) for c in row):
                        report.skipped += 1
                        report.errors.append(f"IDP row {n}: missing Employee ID and Ticket")
                    continue
                if not emp or not ticket:
                    report.skipped += 1
                    report.errors.append(f"IDP row {n}: missing {'Employee ID' if not emp else 'Ticket'}")
                    continue
                name, suffix_dept = N.split_reporter(g("reporter_name"))
                dept = N.to_text(g("department")) or suffix_dept
                sd, ok1 = N.parse_date(g("start_date"))
                ed, ok2 = N.parse_date(g("end_date"))
                for ok, col in ((ok1, "Course Start Date"), (ok2, "Course End Date")):
                    if not ok:
                        report.errors.append(f"IDP row {n}: unparsable date in column '{col}'")
                if (emp, ticket) in seen:
                    report.skipped += 1
                    continue
                seen.add((emp, ticket))
                out.append(IdpRow(
                    emp_no=emp, reporter_name=name, job_level=N.to_text(g("job_level")), department=dept,
                    summary=N.to_text(g("summary")), delivery_method=N.to_text(g("delivery_method")),
                    current_status=N.to_text(g("current_status")), idp_type=N.to_text(g("idp_type")),
                    course_category=N.to_text(g("course_category")),
                    attendance_status=N.to_text(g("attendance_status")) or "-",
                    start_date=sd, end_date=ed, ticket=ticket, section_b=N.clean_section_b(g("section_b"))))
            except Exception as exc:  # noqa: BLE001 - one bad row never aborts the whole read
                report.skipped += 1
                report.errors.append(f"IDP row {n}: could not be parsed ({type(exc).__name__})")
        report.errors.extend(tab.notes)
        report.created = len(out)
        report.finished_at = datetime.now(timezone.utc)
        return out, report

    # ---------------------------------------------------------------- employees
    def read_employees(self, path: str) -> list[EmployeeRecord]:
        return self.read_employees_report(path)[0]

    def read_employees_report(self, path: str, sheet: str | None = None) -> tuple[list[EmployeeRecord], SyncReport]:
        wb = _open(path, self.limits)
        try:
            ws = _find_sheet(wb, sheet) if sheet else self._pick(wb, "ifs", IFS_ALIASES, IFS_REQUIRED)
            if ws is None:
                raise ExcelFormatError("IFS sheet not found")
            return self._parse_employees(ws)
        finally:
            wb.close()

    def _parse_employees(self, ws) -> tuple[list[EmployeeRecord], SyncReport]:
        report = SyncReport(source="excel:ifs", started_at=datetime.now(timezone.utc))
        tab = _Table(ws, IFS_ALIASES, IFS_REQUIRED, ws.title, guard=self._guard())
        out: dict[str, EmployeeRecord] = {}
        for n, row in tab.data():
            try:
                g = partial(N.cell, row, tab.cmap)
                emp = N.clean_emp_no(g("emp_no"))
                if not emp:
                    if any(not N.is_blank(c) for c in row):
                        report.skipped += 1
                        report.errors.append(f"IFS row {n}: missing Emp No")
                    continue
                if emp in out:
                    report.skipped += 1
                    report.errors.append(f"IFS row {n}: duplicate Emp No (first occurrence kept)")
                    continue
                out[emp] = EmployeeRecord(
                    emp_no=emp,
                    first_name_en=N.to_text(g("first_en")), last_name_en=N.to_text(g("last_en")),
                    first_name_fa=N.fa_text(g("first_fa")), last_name_fa=N.fa_text(g("last_fa")),
                    gender=N.gender(g("gender")),  # type: ignore[arg-type]
                    email=N.clean_email(g("email")), mobile=N.fix_mobile(g("mobile")),
                    national_code=N.fix_national_code(g("national_code")),
                    division=N.to_text(g("division")) or N.to_text(g("org_name")),
                    cost_center=N.to_text(g("cost_center")), org_code=N.to_text(g("org_code")),
                    position_title=N.to_text(g("position")),
                    supervisor_emp_no=N.clean_emp_no(g("sup_emp_no")),
                    status=N.to_text(g("status")) or "Active")
            except Exception as exc:  # noqa: BLE001 - one bad row never aborts the whole read
                report.skipped += 1
                report.errors.append(f"IFS row {n}: could not be parsed ({type(exc).__name__})")
        report.errors.extend(tab.notes)
        report.created = len(out)
        report.finished_at = datetime.now(timezone.utc)
        return list(out.values()), report

    # ---------------------------------------------------------------- legacy workbook
    def read_legacy(self, path: str) -> LegacyData:
        return self.read_legacy_report(path)[0]

    def read_legacy_report(self, path: str) -> tuple[LegacyData, SyncReport]:
        wb = _open(path, self.limits)
        report = SyncReport(source="excel:legacy", started_at=datetime.now(timezone.utc))
        data = LegacyData()
        try:
            names = {ws.title.strip().lower(): ws for ws in wb.worksheets}
            if "ifs" in names:
                data.employees, r = self._parse_employees(names["ifs"])
                self._merge(report, r)
            if "idp" in names:
                data.idp, r = self._parse_idp(names["idp"])
                self._merge(report, r)
            if "chistory" in names:
                data.runs = self._parse_runs(names["chistory"], report)
            if "record" in names:
                data.records = self._parse_records(names["record"], report)
            if not any(k in names for k in ("ifs", "idp", "chistory", "record")):
                raise ExcelFormatError("None of the legacy sheets (IFS, IDP, Record, CHistory) found")
        finally:
            wb.close()
        report.finished_at = datetime.now(timezone.utc)
        return data, report

    @staticmethod
    def _merge(into: SyncReport, other: SyncReport) -> None:
        into.skipped += other.skipped
        into.errors.extend(other.errors)

    def _parse_runs(self, ws, report: SyncReport) -> list[RunRecord]:
        tab = _Table(ws, RUN_ALIASES, RUN_REQUIRED, ws.title, guard=self._guard())
        out = []
        for n, row in tab.data():
            try:
                g = partial(N.cell, row, tab.cmap)
                name = N.to_text(g("course_name"))
                if not name:
                    continue
                dates = {}
                for key, col in (("start_date", "start Date"), ("end_date", "End Date"),
                                 ("send_date", "Send Date"), ("deadline", "Deadline")):
                    d, ok = N.parse_date(g(key))
                    dates[key] = d
                    if not ok:
                        report.errors.append(f"CHistory row {n}: unparsable date in column '{col}'")
                dur, ok = N.to_float(g("duration"))
                amt, ok2 = N.to_float(g("amount"))
                if not ok or not ok2:
                    report.errors.append(f"CHistory row {n}: non-numeric Duration/Amount")
                typ = N.to_text(g("type")).lower()
                inst = N.to_text(g("institute"))
                out.append(RunRecord(
                    course_name=name, course_name_fa=N.fa_text(g("course_name_fa")),
                    type="online" if "online" in typ else "in_person",
                    institute="" if inst == "0" else inst, duration_h=dur or 0.0,
                    start_date=dates["start_date"], end_date=dates["end_date"],
                    days=N.to_text(g("days")), time=N.to_text(g("time")), sessions=N.to_int(g("sessions")),
                    send_date=dates["send_date"], deadline=dates["deadline"], amount=amt,
                    section_b=N.clean_section_b(g("section_b")), link=N.to_text(g("link"))))
            except Exception as exc:  # noqa: BLE001 - one bad row never aborts the whole read
                report.skipped += 1
                report.errors.append(f"CHistory row {n}: could not be parsed ({type(exc).__name__})")
        report.errors.extend(tab.notes)
        return out

    def _parse_records(self, ws, report: SyncReport) -> list[RecordRow]:
        tab = _Table(ws, RECORD_ALIASES, RECORD_REQUIRED, ws.title, min_score=3, guard=self._guard())
        out = []
        for n, row in tab.data():
            try:
                g = partial(N.cell, row, tab.cmap)
                emp = N.clean_emp_no(g("emp_no"))
                ticket = N.clean_ticket(g("ticket"))
                if not emp and not ticket:
                    continue
                sd, ok = N.parse_date(g("start_date"))
                if not ok:
                    report.errors.append(f"Record row {n}: unparsable date in column 'start Date'")
                inst = N.to_text(g("institute"))
                out.append(RecordRow(
                    emp_no=emp, ticket=ticket, course_name=N.to_text(g("course_name")),
                    institute="" if inst == "0" else inst, start_date=sd,
                    motivation=N.to_text(g("motivation")), registration=N.to_text(g("registration")),
                    calendar=N.to_text(g("calendar")), payment=N.to_text(g("payment")),
                    comment=N.to_text(g("comment"))))
            except Exception as exc:  # noqa: BLE001 - one bad row never aborts the whole read
                report.skipped += 1
                report.errors.append(f"Record row {n}: could not be parsed ({type(exc).__name__})")
        report.errors.extend(tab.notes)
        return out

    # ---------------------------------------------------------------- writer
    def write_legacy(self, data: LegacyData, path: str) -> None:
        """Write sheets IFS, IDP, Record, CHistory in the legacy column layout. The target is written
        atomically (temp file + rename) so a half-written file is never observed."""
        from app.persian import PersianCalendar
        cal = PersianCalendar()
        wb = Workbook(write_only=False)
        wb.remove(wb.active)

        ws = wb.create_sheet("IFS")
        ws.append(LEGACY_IFS_COLS)
        for e in data.employees:
            rb = _RowBuilder(LEGACY_IFS_COLS)
            put = rb.put
            put("Emp No", e.emp_no)
            put("Employee Status", e.status)
            put("Org Code", e.org_code)
            put("Organization Name", e.division)
            put("Cost Center", e.cost_center)
            put("Fname", e.first_name_en)
            put("Lname", e.last_name_en)
            put("Gender", {"male": "Male", "female": "Female"}.get(e.gender, ""))
            put("Mobile", e.mobile)
            put("Email", e.email)
            put("Farsi First Name", e.first_name_fa)
            put("Farsi Last Name", e.last_name_fa)
            put("Position Title", e.position_title)
            put("Sup Emp No", e.supervisor_emp_no)
            put("National Code", e.national_code)
            put("Division", e.division)
            ws.append(rb.values)

        ws = wb.create_sheet("IDP")
        for _ in range(3):
            ws.append([])
        ws.append(LEGACY_IDP_COLS)     # legacy has the raw export header on row 4 and the table header on 5
        ws.append(LEGACY_IDP_COLS)
        for i in data.idp:
            suffix = f" [ MTNIrancell - {i.department} ]" if i.department else ""
            ws.append([i.emp_no, f"{i.reporter_name}{suffix}", i.job_level, i.department, i.summary,
                       i.delivery_method, i.current_status, i.idp_type, i.course_category,
                       i.attendance_status or "-", _dt(i.start_date), _dt(i.end_date), i.ticket,
                       i.section_b or "-"])

        emp_by_no = {e.emp_no: e for e in data.employees}
        run_by_key = {(r.course_name.lower(), r.start_date): r for r in data.runs}
        ws = wb.create_sheet("Record")
        ws.append([])
        ws.append([])
        ws.append(LEGACY_RECORD_COLS)
        ws.append(LEGACY_RECORD_COLS)
        for n, rec in enumerate(data.records, start=1):
            e = emp_by_no.get(rec.emp_no)
            run = run_by_key.get((rec.course_name.lower(), rec.start_date))
            rb = _RowBuilder(LEGACY_RECORD_COLS)
            put = rb.put
            put("Row", n)
            put("Employee Id", rec.emp_no)
            if e:
                put("Fullname", f"{e.first_name_en} {e.last_name_en}".strip())
                put("division", e.division)
                put("Email", e.email)
            put("Ticket", rec.ticket)
            put("Course Name", rec.course_name)
            put("Institute", rec.institute)
            put("start Date", _dt(rec.start_date))
            if rec.start_date:
                put("jStart Date", _short_j(cal, rec.start_date))
            if run:
                put("Course Name2", run.course_name_fa)
                put("Type", "Online" if run.type == "online" else "In-Person")
                put("Duration", run.duration_h)
                put("End Date", _dt(run.end_date))
                put("Time", run.time)
                put("Sessions Count", run.sessions)
                put("Day", run.days)
                put("Send Date", _dt(run.send_date))
                put("Deadline", _dt(run.deadline))
                put("SectionB", run.section_b)
                put("link", run.link)
                put("Amount", run.amount or 0)
            put("Motivation", rec.motivation)
            put("Registration", rec.registration)
            put("Calendar", rec.calendar)
            put("Payment", rec.payment)
            put("Comment", rec.comment)
            ws.append(rb.values)

        ws = wb.create_sheet("CHistory")
        ws.append([])
        ws.append([])
        ws.append(LEGACY_CHISTORY_COLS)
        ws.append(LEGACY_CHISTORY_COLS)
        for n, run in enumerate(data.runs, start=1):
            rb = _RowBuilder(LEGACY_CHISTORY_COLS)
            put = rb.put
            put("Row", n)
            put("Course Name", run.course_name)
            put("Course NameFarsi", run.course_name_fa)
            put("Type", "Online" if run.type == "online" else "In-Person")
            put("Institute", run.institute)
            put("Duration", run.duration_h)
            if run.start_date:
                put("jStart Date", _short_j(cal, run.start_date))
            put("End Date", _dt(run.end_date))
            put("start Date", _dt(run.start_date))
            put("Time", run.time)
            put("Sessions Count", run.sessions)
            put("Day", run.days)
            put("Send Date", _dt(run.send_date))
            put("Deadline", _dl(run.deadline))
            put("SectionB", run.section_b)
            put("Amount", run.amount)
            put("Link", run.link)
            put("Key", f"{run.course_name}{run.institute}{_dt(run.start_date)}")
            ws.append(rb.values)

        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") and v.strip() != "-":
                        c.data_type = "s"        # stays TEXT: never a live formula (spreadsheet injection)
        target = Path(path)
        if target.suffix.lower() != ".xlsx":
            raise ExcelFormatError("write_legacy only writes .xlsx (macro-free); use a .xlsx target")
        tmp = target.with_name(target.name + ".tmp")
        wb.save(tmp)
        tmp.replace(target)


class _RowBuilder:
    def __init__(self, cols: list[str]):
        self._cols = cols
        self.values: list[Any] = [""] * len(cols)

    def put(self, col: str, v: Any) -> None:
        self.values[self._cols.index(col)] = v


def _dt(d: date | None) -> str:
    return d.strftime("%Y/%m/%d") if d else ""


def _dl(d: date | None) -> str:
    return d.strftime("%A-%d-%b-%Y") if d else ""


def _short_j(cal, d: date) -> str:
    y, m, dd = cal.to_jalali(d)
    return f"{y % 100:02d}/{m:02d}/{dd:02d}"     # legacy 'yy/mm/dd' form
