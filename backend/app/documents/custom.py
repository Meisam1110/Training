"""User-created templates: HTML from the rich-text editor -> e-mail HTML body + .docx.

Input is UNTRUSTED. Pipeline: parse with html.parser into a tree, keep only a whitelist of tags and
attributes (script/style/iframe/... are dropped with their content, event handlers and styles are
dropped, only http/https/mailto links survive), then substitute field tokens in ONE pass over the
tokenised text (inserted values are never rescanned, so ``{{`` inside data stays literal). All values
are HTML-escaped in the e-mail body and inserted as plain text in the .docx.

Tokens: ``{{key}}`` (required), ``{{key?}}`` (optional: empty when missing), ``{{key|fa}}`` (Persian
digits), ``{{key|ltr}}``, ``{{key|iso}}``; the editor may also emit ``<span data-field="key">``. ``key``
is a catalog key (``course_name``) or a legacy merge-field name (``Course_Name``).
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from html import escape
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Mapping

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.contracts import RenderedLetter
from app.persian import PersianCalendar, to_latin_digits, to_persian_digits

from . import engine
from .specs import catalog_index

MAX_BODY_CHARS = 500_000
MAX_DEPTH = 40
MAX_TOKENS = 2000
MAX_TABLE_CELLS = 5000

ALLOWED = {"p", "br", "b", "strong", "i", "em", "u", "ul", "ol", "li", "a", "table", "thead", "tbody", "tr", "td", "th",
           "h2", "h3", "span", "div"}
DROP_WITH_CONTENT = {"script", "style", "iframe", "object", "embed", "form", "svg", "math", "noscript", "template",
                     "head", "title", "textarea", "select", "button", "link", "meta", "base", "frame", "frameset"}
VOID = {"br", "hr", "img", "input", "meta", "link", "base", "col", "area", "wbr", "source", "track", "embed"}
TOKEN = re.compile(r"\{\{\s*([A-Za-z_][\w.]*)\s*(\?)?\s*(?:\|\s*(fa|ltr|iso))?\s*\}\}")
FIELD_NAME = re.compile(r"^[A-Za-z_][\w.]*$")
SAFE_SCHEMES = ("http://", "https://", "mailto:")
_CTRL = re.compile(r"[\x00-\x20\x7f-\x9f​-‏ -‮⁦-⁩﻿]")


class UnknownFieldError(ValueError):
    """The template uses field tokens that exist neither in the catalog nor in the supplied fields."""

    def __init__(self, fields: list[str]):
        self.kind, self.fields = "custom", fields
        super().__init__(f"Unknown field(s) in template: {', '.join(fields)}")


class TemplateSyntaxError(ValueError):
    """Body too large / too deeply nested / too many tokens."""


# ------------------------------------------------------------------------------ parsing
class Node:
    __slots__ = ("tag", "attrs", "children")

    def __init__(self, tag: str, attrs: dict[str, str] | None = None):
        self.tag, self.attrs, self.children = tag, attrs or {}, []


class Field:
    """A ``{{key}}`` token produced by ``<span data-field>`` (already normalised)."""
    __slots__ = ("name", "optional", "flt")

    def __init__(self, name: str, optional: bool = False, flt: str | None = None):
        self.name, self.optional, self.flt = name, optional, flt


class _Builder(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]
        self.skip_depth = 0
        self.skip_tag = ""

    def _clean_attrs(self, tag: str, attrs: list[tuple[str, str | None]]) -> dict[str, str]:
        out: dict[str, str] = {}
        for k, v in attrs:
            k = k.lower()
            v = v or ""
            if tag == "a" and k == "href":
                out["href"] = v
            elif tag in ("td", "th") and k in ("colspan", "rowspan") and v.isdigit() and 1 <= int(v) <= 20:
                out[k] = v
            elif tag == "span" and k == "data-field" and FIELD_NAME.match(v):
                out["data-field"] = v
        return out

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if self.skip_depth:
            if tag == self.skip_tag and tag not in VOID:
                self.skip_depth += 1
            return
        if tag in DROP_WITH_CONTENT:
            if tag not in VOID:
                self.skip_depth, self.skip_tag = 1, tag
            return
        if tag not in ALLOWED:
            return                                   # unknown tag: dropped, children kept
        if len(self.stack) > MAX_DEPTH:
            raise TemplateSyntaxError("template is nested too deeply")
        node = Node(tag, self._clean_attrs(tag, attrs))
        self.stack[-1].children.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        if tag.lower() not in VOID and tag.lower() in ALLOWED and not self.skip_depth:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self.skip_depth:
            if tag == self.skip_tag:
                self.skip_depth -= 1
            return
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data: str) -> None:
        if not self.skip_depth and data:
            self.stack[-1].children.append(data)


def parse_html(body_html: str) -> Node:
    if not isinstance(body_html, str):
        raise TypeError("body_html must be str")
    if len(body_html) > MAX_BODY_CHARS:
        raise TemplateSyntaxError(f"template body is too large (max {MAX_BODY_CHARS} characters)")
    b = _Builder()
    b.feed(body_html)
    b.close()
    return b.root


def _tidy_breaks(items: list) -> list:
    """Collapse empty lines created by optional fields that came out empty (``<br>{{x?}}<br>``)."""
    out = []
    for text, fmt in items:
        out.append([text, fmt])
    joined = "".join(t for t, _ in out)
    if "\n" not in joined:
        return [(t, f) for t, f in out]
    res, last_nl = [], True
    for t, f in out:
        buf = []
        for ch in t:
            if ch == "\n":
                if last_nl:
                    continue
                last_nl = True
            elif not ch.isspace():
                last_nl = False
            buf.append(ch)
        res.append(("".join(buf), f))
    while res and res[-1][0].endswith("\n"):
        res[-1] = (res[-1][0].rstrip("\n"), res[-1][1])
    return res


def _split_tokens(text: str) -> list[str | Field]:
    out: list[str | Field] = []
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            out.append(text[pos:m.start()])
        out.append(Field(m.group(1), bool(m.group(2)), m.group(3)))
        pos = m.end()
    if pos < len(text):
        out.append(text[pos:])
    return out


def _expand_fields(node: Node) -> int:
    """Turn data-field spans and text tokens into Field objects (in place). Returns the token count."""
    n = 0
    new: list[Any] = []
    for ch in node.children:
        if isinstance(ch, str):
            parts = _split_tokens(ch)
            n += sum(isinstance(p, Field) for p in parts)
            new.extend(parts)
        elif ch.tag == "span" and "data-field" in ch.attrs:
            new.append(Field(ch.attrs["data-field"]))
            n += 1
        else:
            n += _expand_fields(ch)
            new.append(ch)
    node.children = new
    if n > MAX_TOKENS:
        raise TemplateSyntaxError("template has too many field tokens")
    return n


def _href_tokens(node: Node, found: list[Field]) -> None:
    for ch in node.children:
        if isinstance(ch, Node):
            if ch.tag == "a" and "href" in ch.attrs:
                found.extend(p for p in _split_tokens(ch.attrs["href"]) if isinstance(p, Field))
            _href_tokens(ch, found)


def _walk_fields(node: Node, found: list[Field]) -> None:
    for ch in node.children:
        if isinstance(ch, Field):
            found.append(ch)
        elif isinstance(ch, Node):
            _walk_fields(ch, found)


# ------------------------------------------------------------------------------ values
def _to_date(v: Any) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(to_latin_digits(str(v)).strip()[:10])
    except ValueError:
        return None


class Resolver:
    """Looks tokens up in the supplied fields by catalog key or legacy name, with derivations from raw inputs."""

    def __init__(self, fields: Mapping[str, Any], persian_default: bool, today: date | None = None):
        self.cat = catalog_index()
        self.f: dict[str, Any] = {k: v for k, v in fields.items() if k != "persian_digits"}
        self.persian = bool(fields.get("persian_digits", persian_default))
        cal = PersianCalendar()
        sd, dd = _to_date(self.f.get("start_date")), _to_date(self.f.get("deadline_date"))
        ld = _to_date(self.f.get("letter_date"))
        d = self.f
        if sd:
            d.setdefault("start_date_jalali", cal.format(sd, "short"))
            d.setdefault("start_date_words", cal.format(sd, "words"))
            d.setdefault("start_date_greg", sd.strftime("%Y/%m/%d"))
        if dd:
            jy, jm, jd = cal.to_jalali(dd)
            d.setdefault("deadline_fa", f"{jd} {cal.month_name(jm)} {jy}")
            d.setdefault("deadline_en", f"{dd.strftime('%A-%d-%b-%Y')}")
        if ld:
            d.setdefault("letter_date_jalali", cal.format(ld, "short"))
        hours = d.get("duration_hours")
        if hours not in (None, ""):
            h = f"{float(hours):g}"
            d.setdefault("duration_fa", f"{h} ساعت")
            d.setdefault("duration_en", f"{h} Hours")
        d.setdefault("today_jalali", cal.format(today or date.today(), "short"))
        g = str(d.get("gender", "")).lower()
        if g in ("male", "female"):
            d.setdefault("employee_title_fa", "جناب آقای" if g == "male" else "سرکار خانم")
            d.setdefault("employee_title_en", "Mr." if g == "male" else "Ms.")

    def known(self, name: str) -> bool:
        return name in self.cat or name in self.f

    def raw(self, name: str) -> Any:
        if name in self.f and self.f[name] not in (None, ""):
            return self.f[name]
        e = self.cat.get(name)
        if e:
            for cand in [e["key"], *e["legacy"]]:
                v = self.f.get(cand)
                if v not in (None, ""):
                    return v
        return None

    def text(self, name: str, flt: str | None) -> str | None:
        v = self.raw(name)
        if v is None or isinstance(v, (list, tuple)):
            return None if v is None else ""
        if isinstance(v, (date, datetime)):
            s = (v.date() if isinstance(v, datetime) else v).isoformat()
        else:
            s = str(v).strip()
        s = engine.clean_text(s)
        e = self.cat.get(name)
        typ = e["type"] if e else "text"
        if typ in ("date_jalali", "number", "money") and self.persian and flt is None:
            s = to_persian_digits(s)
        return engine.apply_filter(s, flt, self.persian)


# ------------------------------------------------------------------------------ HTML output
FONT = "Tahoma,Arial,sans-serif"
YELLOW, BLUE, GREY = "#FFCB05", "#004A8F", "#BFBFBF"


def _safe_href(url: str) -> str | None:
    u = _CTRL.sub("", url or "")
    return u if u.lower().startswith(SAFE_SCHEMES) and len(u) < 2000 else None


def _list_table(items: list[Mapping[str, Any]], list_name: str = "", cat: Mapping[str, dict] | None = None) -> Node:
    """A list-valued token becomes a plain table (header = item keys)."""
    t = Node("table")
    keys = list(items[0].keys()) if items else []
    head = Node("tr")
    for k in keys:
        th = Node("th")
        label = (cat or {}).get(f"{list_name}.{k}", {}).get("label_fa")
        th.children.append(label or str(k))
        head.children.append(th)
    t.children.append(head)
    for it in items:
        tr = Node("tr")
        for k in keys:
            td = Node("td")
            td.children.append(engine.clean_text(str(it.get(k, ""))))
            tr.children.append(td)
        t.children.append(tr)
    return t


class _Out:
    def __init__(self, res: Resolver, rtl: bool):
        self.res, self.rtl = res, rtl
        self.values: dict[str, str] = {}
        self.cells = 0

    # field -> plain string
    def field_text(self, f: Field) -> str:
        v = self.res.text(f.name, f.flt) or ""
        self.values[f.name] = v
        return v

    def sub_attr(self, s: str) -> str:
        return "".join(p if isinstance(p, str) else self.field_text(p) for p in _split_tokens(s))

    def expand_lists(self, node: Node) -> None:
        """A paragraph holding only a list-valued token turns into a table node."""
        new = []
        for ch in node.children:
            if isinstance(ch, Node):
                only = [c for c in ch.children if not (isinstance(c, str) and not c.strip())]
                if ch.tag in ("p", "div") and len(only) == 1 and isinstance(only[0], Field):
                    val = self.res.raw(only[0].name)
                    if isinstance(val, (list, tuple)) and val and all(isinstance(x, Mapping) for x in val):
                        new.append(_list_table([dict(x) for x in val], only[0].name, self.res.cat))
                        continue
                self.expand_lists(ch)
            new.append(ch)
        node.children = new

    # ---- html
    def html(self, node: Node) -> str:
        d, al = ("rtl", "right") if self.rtl else ("ltr", "left")
        out = []
        for ch in node.children:
            if isinstance(ch, str):
                out.append(escape(engine.clean_text(ch), quote=False))
            elif isinstance(ch, Field):
                out.append(escape(self.field_text(ch), quote=False))
            else:
                out.append(self._el(ch, d, al))
        return "".join(out)

    def _el(self, n: Node, d: str, al: str) -> str:
        t = n.tag
        inner = self.html(n)
        base = f"font-family:{FONT};font-size:14px;line-height:1.7;color:#222222;direction:{d};text-align:{al};"
        if t == "br":
            return "<br>"
        if t in ("p", "div"):
            return f'<p dir="{d}" style="{base}margin:0 0 12px 0;">{inner}</p>'
        if t == "h2":
            return f'<h2 dir="{d}" style="{base}font-size:18px;font-weight:bold;color:{BLUE};margin:0 0 10px 0;">{inner}</h2>'
        if t == "h3":
            return f'<h3 dir="{d}" style="{base}font-size:16px;font-weight:bold;color:{BLUE};margin:0 0 8px 0;">{inner}</h3>'
        if t in ("b", "strong"):
            return f"<b>{inner}</b>"
        if t in ("i", "em"):
            return f"<i>{inner}</i>"
        if t == "u":
            return f"<u>{inner}</u>"
        if t in ("ul", "ol"):
            return f'<{t} dir="{d}" style="{base}margin:0 0 12px 0;padding:0 22px 0 22px;">{inner}</{t}>'
        if t == "li":
            return f'<li style="margin:0 0 6px 0;">{inner}</li>'
        if t == "a":
            href = _safe_href(self.sub_attr(n.attrs.get("href", "")))
            if href:
                return f'<a href="{escape(href, quote=True)}" style="color:#0563C1;text-decoration:underline;">{inner}</a>'
            return inner
        if t == "table":
            return (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" border="0" dir="{d}" '
                    f'style="border-collapse:collapse;margin:0 0 12px 0;">{inner}</table>')
        if t in ("thead", "tbody"):
            return inner
        if t == "tr":
            return f"<tr>{inner}</tr>"
        if t in ("td", "th"):
            self.cells += 1
            if self.cells > MAX_TABLE_CELLS:
                raise TemplateSyntaxError("template table is too large")
            attrs = "".join(f' {k}="{v}"' for k, v in n.attrs.items() if k in ("colspan", "rowspan"))
            bg = f' bgcolor="{YELLOW}"' if t == "th" else ""
            w = "font-weight:bold;" if t == "th" else ""
            return (f'<{t}{attrs}{bg} dir="{d}" align="{al}" style="{base}{w}border:1px solid {GREY};padding:6px 10px;">'
                    f"{inner}</{t}>")
        return inner                                  # span and anything else: content only

    # ---- docx
    def to_docx(self, doc: Document, node: Node, rtl: bool) -> None:
        for ch in node.children:
            if isinstance(ch, (str, Field)):
                txt = ch if isinstance(ch, str) else self.field_text(ch)
                if txt.strip():
                    self._para(doc, [(txt, {})], rtl)
            else:
                self._block(doc, ch, rtl)

    def _runs(self, node: Node, fmt: dict, acc: list) -> None:
        for ch in node.children:
            if isinstance(ch, str):
                acc.append((engine.clean_text(ch), dict(fmt)))
            elif isinstance(ch, Field):
                acc.append((self.field_text(ch), dict(fmt)))
            elif ch.tag == "br":
                acc.append(("\n", dict(fmt)))
            elif ch.tag in ("b", "strong"):
                self._runs(ch, {**fmt, "b": True}, acc)
            elif ch.tag in ("i", "em"):
                self._runs(ch, {**fmt, "i": True}, acc)
            elif ch.tag == "u":
                self._runs(ch, {**fmt, "u": True}, acc)
            elif ch.tag == "a":
                href = _safe_href(self.sub_attr(ch.attrs.get("href", "")))
                self._runs(ch, {**fmt, "link": href} if href else fmt, acc)
            elif ch.tag in ("ul", "ol", "table", "p", "div", "h2", "h3"):
                continue
            else:
                self._runs(ch, fmt, acc)

    def _para(self, container, runs: list, rtl: bool, *, size: float | None = None, bold=False, color=None,
              prefix: str = "", space_after: int = 6):
        p = container.add_paragraph()
        p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(space_after), Pt(0)
        if rtl:
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        items = ([(prefix, {})] if prefix else []) + list(runs)
        items = _tidy_breaks(items)
        for text, fmt in items:
            text = engine.clean_text(text)
            if not text:
                continue
            link = fmt.get("link")
            if link:
                self._hyperlink(p, text, link, rtl)
                continue
            parts = text.split("\n")
            for i, part in enumerate(parts):
                r = p.add_run(part)
                r.bold = (bold or fmt.get("b")) or None
                r.italic = fmt.get("i") or None
                r.underline = fmt.get("u") or None
                if size:
                    r.font.size = Pt(size)
                if color is not None:
                    r.font.color.rgb = color
                if rtl:
                    rpr = r._r.get_or_add_rPr()
                    rpr.append(OxmlElement("w:rtl"))
                if i < len(parts) - 1:
                    r.add_break()
        return p

    def _hyperlink(self, p, text: str, url: str, rtl: bool) -> None:
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "0563C1")
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(c)
        rpr.append(u)
        if rtl:
            rpr.append(OxmlElement("w:rtl"))
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    def _block(self, doc, n: Node, rtl: bool, depth: int = 0) -> None:
        t = n.tag
        if t in ("p", "div", "h2", "h3"):
            runs: list = []
            self._runs(n, {}, runs)
            size = {"h2": 15, "h3": 13}.get(t)
            if any(r[0].strip() for r in runs):
                self._para(doc, runs, rtl, size=size, bold=t in ("h2", "h3"), color=BLUE_RGB if t in ("h2", "h3") else None)
            for ch in n.children:                     # nested blocks inside a div
                if isinstance(ch, Node) and ch.tag in ("ul", "ol", "table", "p", "div", "h2", "h3"):
                    self._block(doc, ch, rtl, depth + 1)
        elif t in ("ul", "ol"):
            for i, li in enumerate([c for c in n.children if isinstance(c, Node) and c.tag == "li"], start=1):
                runs = []
                self._runs(li, {}, runs)
                self._para(doc, runs, rtl, prefix=("•  " if t == "ul" else f"{i}.  "), space_after=3)
        elif t == "table":
            self._table(doc, n, rtl)
        elif t in ("thead", "tbody", "span", "a", "b", "strong", "i", "em", "u", "li"):
            runs = []
            wrapper = Node("x")
            wrapper.children = [n]
            self._runs(wrapper, {}, runs)
            if any(r[0].strip() for r in runs):
                self._para(doc, runs, rtl)

    def _table(self, doc, n: Node, rtl: bool) -> None:
        rows = []
        for ch in n.children:
            if isinstance(ch, Node) and ch.tag in ("thead", "tbody"):
                rows += [c for c in ch.children if isinstance(c, Node) and c.tag == "tr"]
            elif isinstance(ch, Node) and ch.tag == "tr":
                rows.append(ch)
        if not rows:
            return
        ncols = max(sum(1 for c in r.children if isinstance(c, Node) and c.tag in ("td", "th")) for r in rows) or 1
        t = doc.add_table(rows=0, cols=ncols)
        tblpr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "BFBFBF")
            b.append(e)
        tblpr.append(b)
        if rtl:
            tblpr.append(OxmlElement("w:bidiVisual"))
        for r in rows:
            cells = [c for c in r.children if isinstance(c, Node) and c.tag in ("td", "th")]
            self.cells += len(cells)
            if self.cells > MAX_TABLE_CELLS:
                raise TemplateSyntaxError("template table is too large")
            row = t.add_row()
            for i, c in enumerate(cells[:ncols]):
                runs = []
                self._runs(c, {}, runs)
                cell = row.cells[i]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if rtl:
                    p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
                text = "".join(x[0] for x in runs)
                run = p.add_run(text)
                run.font.size = Pt(9.5)
                if c.tag == "th":
                    run.bold = True
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "FFCB05")
                    cell._tc.get_or_add_tcPr().append(shd)
                if rtl:
                    run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
        doc.add_paragraph().paragraph_format.space_after = Pt(4)


BLUE_RGB = RGBColor(0x00, 0x4A, 0x8F)


def _new_document(templates_dir: Path, letterhead: bool):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Tahoma", Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Tahoma")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    top, bottom = templates_dir / "assets" / "letterhead_top.jpg", templates_dir / "assets" / "letterhead_bottom.jpg"
    if letterhead and top.is_file() and bottom.is_file():
        sec.header_distance, sec.footer_distance = Cm(0.5), Cm(0.5)
        sec.top_margin, sec.bottom_margin = Cm(4.7), Cm(2.4)
        for part, img in ((sec.header, top), (sec.footer, bottom)):
            p = part.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(img), width=Cm(19.0))
    cp = doc.core_properties
    cp.author = cp.last_modified_by = "TMS"
    cp.created = cp.modified = datetime(2025, 1, 1)
    return doc


def render_custom_letter(templates_dir: Path, persian_default: bool, subject: str, body_html: str,
                         fields: Mapping[str, Any], lang: str = "fa", letterhead: bool = True) -> RenderedLetter:
    if lang not in ("fa", "en"):
        raise ValueError("lang must be 'fa' or 'en'")
    if not isinstance(subject, str) or len(subject) > 1000:
        raise TemplateSyntaxError("subject must be a string of at most 1000 characters")
    root = parse_html(body_html)
    _expand_fields(root)
    res = Resolver(fields, persian_default)

    tokens: list[Field] = []
    _walk_fields(root, tokens)
    _href_tokens(root, tokens)
    tokens += [p for p in _split_tokens(subject) if isinstance(p, Field)]
    unknown = sorted({t.name for t in tokens if not res.known(t.name)})
    if unknown:
        raise UnknownFieldError(unknown)
    from .renderer import MissingFieldError
    missing = sorted({t.name for t in tokens if not t.optional and res.raw(t.name) is None})
    if missing:
        raise MissingFieldError("custom", missing)

    rtl = lang == "fa"
    out = _Out(res, rtl)
    out.expand_lists(root)
    inner = re.sub(r"(<br>\s*){2,}", "<br>", out.html(root))
    inner = re.sub(r"<br>\s*(</p>)", r"\1", inner)
    d, al = ("rtl", "right") if rtl else ("ltr", "left")
    subj = "".join(p if isinstance(p, str) else (res.text(p.name, p.flt) or "") for p in _split_tokens(subject))
    subj = " ".join(engine.clean_text(subj).split())
    html = ('<!DOCTYPE html><html lang="%s" dir="%s"><head><meta charset="utf-8"><meta name="viewport" '
            'content="width=device-width, initial-scale=1"><title>%s</title></head>'
            '<body style="margin:0;padding:0;background-color:#F4F4F4;">'
            '<table role="presentation" width="100%%" cellpadding="0" cellspacing="0" border="0" bgcolor="#F4F4F4">'
            '<tr><td align="center" style="padding:16px 8px;">'
            '<!--[if mso]><table role="presentation" width="640" align="center" cellpadding="0" cellspacing="0" '
            'border="0"><tr><td><![endif]-->'
            '<div style="max-width:640px;margin:0 auto;"><table role="presentation" width="100%%" cellpadding="0" '
            'cellspacing="0" border="0" bgcolor="#FFFFFF" style="width:100%%;background-color:#FFFFFF;">'
            '<tr><td height="6" bgcolor="%s" style="height:6px;line-height:6px;font-size:0;">&nbsp;</td></tr>'
            '<tr><td dir="%s" align="%s" style="padding:16px 24px;word-break:break-word;overflow-wrap:anywhere;">%s</td></tr>'
            '</table></div><!--[if mso]></td></tr></table><![endif]--></td></tr></table></body></html>'
            ) % (lang, d, escape(subj), YELLOW, d, al, inner)

    doc = _new_document(Path(templates_dir), letterhead)
    out2 = _Out(res, rtl)
    out2.to_docx(doc, root, rtl)
    buf = io.BytesIO()
    doc.save(buf)
    values = dict(out.values)
    return RenderedLetter(kind="custom", filename="custom_letter.docx", docx_bytes=buf.getvalue(),  # type: ignore[arg-type]
                          html_body=html, subject=subj, fields=values)
