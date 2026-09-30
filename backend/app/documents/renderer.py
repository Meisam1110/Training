"""LetterRenderer: renders the six letters from templates/*.docx (+ an HTML e-mail body).

Strict: a missing/blank required field raises ``MissingFieldError`` naming the field(s); there are no
silent blanks (unlike the legacy mail-merge). Raw dates (``start_date``, ``deadline_date``,
``letter_date``), ``duration_hours`` and ``gender`` may be supplied instead of the derived text fields.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal, Mapping, Sequence

from docx import Document

from app.automation.batching import split_invoice
from app.contracts import LetterKind, RenderedLetter
from app.persian import PersianCalendar, to_latin_digits

from . import engine
from .html import build_html
from .specs import SPECS, F, KindSpec
from .wording import DOCS, SUBJECTS

EN_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
EN_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MONEY_FIELDS = {"invoice_total", "per_attendee", "amount"}
RESERVED = {"persian_digits"}
RAW = {"start_date", "deadline_date", "letter_date", "duration_hours", "gender", "manager_gender",
       "reporter_gender", "unit_amount"}


class MissingFieldError(ValueError):
    """Required field(s) missing or blank. ``fields`` lists canonical names (list items as ``list[3].field``)."""

    def __init__(self, kind: str, fields: list[str]):
        self.kind, self.fields = kind, fields
        super().__init__(f"Cannot render '{kind}': missing required field(s): {', '.join(fields)}")


class TemplateError(RuntimeError):
    pass


def _as_date(v: Any) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(to_latin_digits(str(v)).strip()[:10])
    except ValueError:
        raise ValueError(f"invalid date value: {str(v)[:20]!r} (expected ISO yyyy-mm-dd)") from None


def _num(x: Any) -> str:
    try:
        d = Decimal(to_latin_digits(str(x)).replace(",", "").strip())
    except InvalidOperation:
        return str(x)
    return str(int(d)) if d == d.to_integral() else format(d.normalize(), "f")


def _money(x: Any) -> str:
    try:
        d = Decimal(to_latin_digits(str(x)).replace(",", "").strip())
    except InvalidOperation:
        return str(x)
    return f"{int(d):,}" if d == d.to_integral() else f"{d:,.2f}"


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


class DocxLetterRenderer:
    def __init__(self, templates_dir: str, calendar: PersianCalendar | None = None, persian_digits: bool = True):
        self.dir = Path(templates_dir)
        self.cal = calendar or PersianCalendar()
        self.persian_digits = persian_digits
        self._cache: dict[str, bytes] = {}
        self.defaults = self._load_defaults()

    def _load_defaults(self) -> dict[str, str]:
        """Optional ``letter_defaults.json`` in the templates dir: organisation-wide constants (company
        economic code, national ID, postal code, address, standard signer...) an admin sets once.
        Values given per request always win. The repo ships only ``letter_defaults.example.json``."""
        import json
        p = self.dir / "letter_defaults.json"
        if not p.is_file():
            return {}
        data = json.loads(p.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise TemplateError("letter_defaults.json must be a JSON object")
        known = {f.name for sp in SPECS.values() for f in sp.fields}
        bad = sorted(k for k in data if k not in known)
        if bad:
            raise TemplateError(f"letter_defaults.json: unknown field(s): {', '.join(bad)}")
        return {k: str(v) for k, v in data.items()}

    # ------------------------------------------------------------------ public API
    def required_fields(self, kind: LetterKind, lang: Literal["fa", "en", "both"] = "both") -> Sequence[str]:
        spec = self._spec(kind)
        # Fields that can arrive under an alternative spelling are validated by render() (after alias
        # normalisation) instead of being listed here, so a caller doing `k not in fields` on this list
        # works with either spelling. render() still names them if absent.
        aliased = set(spec.aliases.values())
        names = [f.name for f in spec.fields if f.required and self._in_lang(spec, f, lang)
                 and f.name not in self.defaults and f.name not in aliased]
        names += [ls.name for ls in spec.lists if ls.min_items > 0 and ls.name not in aliased]
        return names

    def render(self, kind: LetterKind, fields: Mapping[str, Any],
               lang: Literal["fa", "en", "both"] = "both") -> RenderedLetter:
        spec = self._spec(kind)
        if lang not in ("fa", "en", "both"):
            raise ValueError("lang must be 'fa', 'en' or 'both'")
        eff = lang if spec.bilingual else "both"
        persian = bool(fields.get("persian_digits", self.persian_digits))
        raw = self._prepare(spec, fields)
        values, lists = self._validate_and_format(spec, raw, eff)

        doc = Document(io.BytesIO(self._template_bytes(spec)))
        unknown = engine.template_placeholders(doc) - self._allowed_names(spec)
        if unknown:
            raise TemplateError(f"template {spec.template} uses undeclared placeholder(s): {', '.join(sorted(unknown))}")
        engine.filter_language(doc, eff)

        def fmt(v: str, flt: str | None) -> str:
            return engine.shield(engine.apply_filter(v, flt, persian))

        def scalar(name: str, flt: str | None) -> str | None:
            if "." in name or name not in values:
                return None
            return fmt(values[name], flt)

        def item_res(lname: str, item: Mapping[str, str], idx: int):
            def r(name: str, flt: str | None) -> str | None:
                if "." in name:
                    ln, fld = name.split(".", 1)
                    if ln != lname:
                        return None
                    return fmt(str(idx) if fld == "_n" else item.get(fld, ""), flt)
                return scalar(name, flt)
            return r

        hide = frozenset(f.name for f in spec.fields if f.hide_if_empty)
        engine.expand_rows(doc, lists, item_res)
        try:
            engine.replace_everywhere(doc, scalar, hide)
        except KeyError as exc:
            raise TemplateError(f"placeholder {exc.args[0]} has no value") from None
        left = engine.leftover_placeholders(doc)
        if left:
            raise TemplateError(f"unresolved placeholders after rendering: {', '.join(sorted(set(left)))}")
        engine.unshield_document(doc)
        buf = io.BytesIO()
        doc.save(buf)

        sections = [(lg, blocks) for lg, blocks in DOCS[kind].items() if eff in ("both", lg)]
        title = self._subject(kind, values, eff)
        html = build_html(sections, values, lists, persian, title)
        flat: dict[str, Any] = dict(values)
        flat.update({k: list(v) for k, v in lists.items()})
        return RenderedLetter(kind=kind, filename=self._filename(kind, values, eff), docx_bytes=buf.getvalue(),
                              html_body=html, subject=title, fields=flat)

    # ------------------------------------------------------------------ internals
    @staticmethod
    def _spec(kind: str) -> KindSpec:
        if kind not in SPECS:
            raise ValueError(f"unknown letter kind: {kind}")
        return SPECS[kind]

    @staticmethod
    def _in_lang(spec: KindSpec, f: F, lang: str) -> bool:
        return not spec.bilingual or lang == "both" or f.lang in (None, lang)

    def _template_bytes(self, spec: KindSpec) -> bytes:
        if spec.template not in self._cache:
            p = self.dir / spec.template
            if not p.is_file():
                raise FileNotFoundError(f"template file not found: {spec.template}")
            self._cache[spec.template] = p.read_bytes()
        return self._cache[spec.template]

    @staticmethod
    def _allowed_names(spec: KindSpec) -> set[str]:
        names = {f.name for f in spec.fields} | {"guide_link"}
        for ls in spec.lists:
            names |= {f"{ls.name}.{f.name}" for f in ls.item_fields} | {f"{ls.name}._n"}
        return names

    def _jal(self, d: date) -> str:
        return self.cal.format(d, "short")

    def _prepare(self, spec: KindSpec, fields: Mapping[str, Any]) -> dict[str, Any]:
        f: dict[str, Any] = {k: v for k, v in fields.items() if k not in RESERVED}
        for alias, canon in spec.aliases.items():
            if alias in f and _blank(f.get(canon)):
                f[canon] = f[alias]
        for ls in spec.lists:
            items = f.get(ls.name)
            if isinstance(items, (list, tuple)) and ls.item_aliases:
                f[ls.name] = [self._alias_item(it, ls.item_aliases) for it in items]
        kind = spec.kind
        for k, v in self.defaults.items():
            if k in {x.name for x in spec.fields} and _blank(f.get(k)):
                f[k] = v

        def setdef(name: str, value: Any) -> None:
            if _blank(f.get(name)) and value is not None:
                f[name] = value

        sd, dd, ld = _as_date(f.get("start_date")), _as_date(f.get("deadline_date")), _as_date(f.get("letter_date"))
        hours = f.get("duration_hours")
        if kind == "course_announcement":
            if sd:
                setdef("jStart_Date", self._jal(sd))
                setdef("Date_Horoof", self.cal.format(sd, "words"))
                setdef("start_Date", sd.strftime("%Y/%m/%d"))
            if dd:
                jy, jm, jd = self.cal.to_jalali(dd)
                setdef("deadline_farsi", f"{jd} {self.cal.month_name(jm)} {jy}")
                setdef("Deadline", f"{EN_DAYS[dd.weekday()]}-{dd.day:02d}-{EN_MONTHS[dd.month - 1]}-{dd.year}")
            if not _blank(hours):
                setdef("Time_Farsi", f"{_num(hours)} ساعت")
                setdef("Duration", f"{_num(hours)} Hours")
            setdef("guide_link", f.get("link"))
        elif kind == "institute_letter":
            if ld:
                setdef("date", self._jal(ld))
            if sd:
                setdef("Jstart_Date", self._jal(sd))
            if not _blank(hours):
                setdef("Duration", f"{_num(hours)} ساعت")
        elif kind == "settlement_letter":
            if sd:
                setdef("start_date_gregorian", sd.strftime("%Y/%m/%d"))
                setdef("start_date_jalali", self._jal(sd))
            att = f.get("attendees")
            if isinstance(att, (list, tuple)):
                setdef("attendee_count", len(att))
                shares: list[Decimal] = []
                if not _blank(f.get("invoice_total")) and len(att) > 0:
                    try:
                        total = Decimal(to_latin_digits(str(f["invoice_total"])).replace(",", ""))
                        shares = split_invoice(total, len(att), places=0 if total == total.to_integral() else 2)
                        setdef("per_attendee", shares[0])
                    except (InvalidOperation, ValueError):
                        shares = []
                items = []
                for n, it in enumerate(att):
                    it = dict(it)
                    if _blank(it.get("label")) and not _blank(it.get("full_name")):
                        it["label"] = f"{f.get('purpose', '')}-{f.get('mtc_no', '')}-{it['full_name']}"
                    if _blank(it.get("amount")):
                        it["amount"] = shares[n] if shares else "-"
                    items.append(it)
                f["attendees"] = items
        elif kind == "level3_request":
            g = str(f.get("gender", "")).lower()
            if g in ("male", "female"):
                setdef("Ftitle", "جناب آقای" if g == "male" else "سرکار خانم")
                setdef("En_title", "Mr." if g == "male" else "Ms.")
        elif kind == "level3_manager_check":
            for prefix, fa_key, en_key in (("manager", "Manager_fTitle", "Manager_Title"),
                                           ("reporter", "Reporter_F_tilte", "Reporter_title")):
                g = str(f.get(f"{prefix}_gender", "")).lower()
                if g in ("male", "female"):
                    setdef(fa_key, "جناب آقای" if g == "male" else "سرکار خانم")
                    setdef(en_key, "Mr." if g == "male" else "Ms.")
        elif kind == "line_manager_approval":
            if "employees" not in f:      # legacy Empname1..9 / EmpSummary1..9 / EmpTicket1..9
                legacy = []
                for i in range(1, 10):
                    if not _blank(f.get(f"Empname{i}")):
                        legacy.append({"name": f.get(f"Empname{i}"), "summary": f.get(f"EmpSummary{i}", ""),
                                       "ticket": f.get(f"EmpTicket{i}", "")})
                if legacy:
                    f["employees"] = legacy
        return f

    @staticmethod
    def _alias_item(item: Any, aliases: Mapping[str, str]) -> Any:
        if not isinstance(item, Mapping):
            return item
        out = dict(item)
        for alias, canon in aliases.items():
            if alias in out and _blank(out.get(canon)):
                out[canon] = out[alias]
        return out

    def _validate_and_format(self, spec: KindSpec, f: dict[str, Any], lang: str):
        missing: list[str] = []
        values: dict[str, str] = {}
        for fs in spec.fields:
            v = f.get(fs.name)
            if _blank(v):
                if fs.required and self._in_lang(spec, fs, lang):
                    missing.append(fs.name)
                    continue
                v = fs.default if fs.default is not None else ""
            values[fs.name] = self._fmt(fs.name, v)
        lists: dict[str, list[dict[str, str]]] = {}
        for ls in spec.lists:
            items = f.get(ls.name)
            if _blank(items) if not isinstance(items, (list, tuple)) else False:
                items = []
            if not isinstance(items, (list, tuple)):
                raise MissingFieldError(spec.kind, [f"{ls.name} (must be a list)"])
            if len(items) < ls.min_items:
                missing.append(ls.name)
            if ls.max_items and len(items) > ls.max_items:
                raise ValueError(f"'{ls.name}' has {len(items)} items; at most {ls.max_items} fit in one "
                                 f"{spec.kind} letter - split them with automation.batch_by_manager")
            out = []
            for n, it in enumerate(items, start=1):
                row: dict[str, str] = {}
                for fs in ls.item_fields:
                    v = it.get(fs.name) if isinstance(it, Mapping) else None
                    if _blank(v):
                        if fs.required:
                            missing.append(f"{ls.name}[{n}].{fs.name}")
                        v = ""
                    row[fs.name] = self._fmt(fs.name, v)
                out.append(row)
            lists[ls.name] = out
        if missing:
            raise MissingFieldError(spec.kind, missing)
        return values, lists

    @staticmethod
    def _fmt(name: str, v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        if name in MONEY_FIELDS:
            return _money(v)
        if isinstance(v, (int, float, Decimal)) and not isinstance(v, bool):
            return _num(v)
        return str(v).strip()

    def _subject(self, kind: str, values: Mapping[str, str], lang: str) -> str:
        subj = SUBJECTS[kind]
        parts = [t for lg, t in subj.items() if lang in ("both", lg)]
        out = " / ".join(re.sub(r"\{\{\s*(\w+)\s*\}\}", lambda m: values.get(m.group(1), ""), t) for t in parts)
        return " ".join(engine.clean_text(out).split())        # single line: no CR/LF ever reaches a header

    @staticmethod
    def _filename(kind: str, values: Mapping[str, str], lang: str) -> str:
        if kind == "course_announcement":
            return f"course_announcement_{lang}.docx"
        if kind == "institute_letter":
            tag = re.sub(r"[^A-Za-z0-9_-]+", "-", to_latin_digits(values.get("Letter_No", ""))).strip("-")
            return f"institute_letter_{tag}.docx" if tag else "institute_letter.docx"
        if kind in ("level3_request", "level3_manager_check") and lang != "both":
            return f"{kind}_{lang}.docx"
        return f"{kind}.docx"
