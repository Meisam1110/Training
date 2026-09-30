"""Regression tests for the independent QA findings owned by integrations (docs/QA_REPORT.md)."""
import io
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal

import httpx
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.automation import compute_due
from app.contracts import EnrollmentFact, OutgoingEmail
from app.documents import DocxLetterRenderer, engine
from app.integrations import factory
from app.integrations.email import (
    GraphSender,
    IdempotencyConflict,
    IdempotencyInProgress,
    InMemoryIdempotencyStore,
    MessageValidationError,
    SmtpSender,
)
from app.integrations.email.common import hash_key
from app.integrations.excel import ExcelFormatError, OpenpyxlExcelGateway
from app.integrations.excel.gateway import ExcelLimits, ExcelReadAborted, check_zip_bomb
from app.integrations.ifs import IfsConfigError, IfsEmployeeSource, IfsUpstreamError
from app.integrations.ifs.mock_server import CLIENT_ID, CLIENT_SECRET, MockIfsServer
from app.integrations.urlcheck import UnsafeUrlError, check_url

from .samples import SAMPLES, TEMPLATES_DIR
from .synth import idp_row, ifs_row, make_idp, make_ifs
from .test_email import Handler, authenticator, free_port, mail  # noqa: F401
from .test_email import cfg as smtp_cfg

Controller = __import__("aiosmtpd.controller", fromlist=["Controller"]).Controller


def docx_text(b: bytes) -> str:
    d = Document(io.BytesIO(b))
    return "\n".join(engine.paragraph_text(p) for root in engine.story_roots(d) for p in engine.all_paragraphs(root))


# ============================================================= QA-B1: Core's real payload shape
def core_settlement_fields(n=5, total="202,500,000", supplier_code="SUP-9"):
    """Exactly the keys/values api/services/letters.py::_run_plan builds for settlement_letter."""
    return {
        "addressee": "Finance Team", "supplier": "Example Institute", "invoice_no": "INV-1",
        "purpose": "Example Course", "mtc_no": "MTC-000001", "start_date_gregorian": "2027/03/01",
        "start_date_jalali": "1405/12/10", "attendee_count": str(n), "supplier_code": supplier_code,
        "invoice_total": total, "per_attendee": "40,500,000",
        "attendees": [{"staff_id": f"EMP{i:04d}", "label": f"Example Course-MTC-000001-Test Person{i}",
                       "full_name": f"Test Person{i}", "division_cost_center": "NWG-23700"} for i in range(n)],
        "sender_name": "Talent Team",
    }


def real_check_required(renderer, kind, fields, lang="both"):
    try:
        from app.api.services.letters import check_required
    except Exception:  # Core not importable in this environment: replicate its 3-line logic
        def check_required(r, k, f, lg="both"):
            missing = [x for x in r.required_fields(k, lg) if x not in f]
            if missing:
                raise ValueError(missing)
    check_required(renderer, kind, fields, lang)


def test_b1_core_payload_through_real_factory():
    r = factory.build_letter_renderer(TEMPLATES_DIR)                      # real factory, real templates
    f = core_settlement_fields()
    real_check_required(r, "settlement_letter", f)                       # Core's own pre-check passes
    out = r.render("settlement_letter", f, "both")
    t = docx_text(out.docx_bytes)
    assert "202,500,000" in t and "Test Person4" in t and "2027/03/01 (1405/12/10)" in t
    assert out.filename == "settlement_letter.docx" and "INV-1" in out.subject


def test_b1_new_spelling_and_mixed_spelling_also_work():
    r = factory.build_letter_renderer(TEMPLATES_DIR)
    core = core_settlement_fields()
    new = {k: v for k, v in core.items() if k not in ("start_date_gregorian", "invoice_total", "attendees")}
    new.update(start_date_greg=core["start_date_gregorian"], total=core["invoice_total"],
               rows=[{"staff_id": a["staff_id"], "description": a["label"], "division_cost_center": "X-1"}
                     for a in core["attendees"]])
    real_check_required(r, "settlement_letter", new)                     # required_fields() is alias-tolerant
    out = r.render("settlement_letter", new)
    assert out.fields["invoice_total"] == "202,500,000" and out.fields["attendees"][0]["label"].endswith("Person0")
    mixed = dict(core, total="1,000")
    del mixed["invoice_total"]
    assert r.render("settlement_letter", mixed).fields["invoice_total"] == "1,000"
    with pytest.raises(Exception, match="invoice_total"):
        r.render("settlement_letter", {k: v for k, v in core.items() if k != "invoice_total"})


def test_b1_empty_supplier_code_and_exact_row_amounts():
    r = DocxLetterRenderer(TEMPLATES_DIR)
    f = core_settlement_fields(n=3, total="100", supplier_code="")
    f["per_attendee"] = "34"
    out = r.render("settlement_letter", f)
    assert out.fields["supplier_code"] == "-"
    amounts = [Decimal(a["amount"].replace(",", "")) for a in out.fields["attendees"]]
    assert amounts == [Decimal(34), Decimal(33), Decimal(33)] and sum(amounts) == 100     # per-row exact shares
    assert "Amount" in docx_text(out.docx_bytes) and "base share" in docx_text(out.docx_bytes)


# ============================================================= QA-M5: no template injection
HOSTILE = ["{{Letter_No}}", "{{participants.email}}", "{{x}}", "{%- if 1 -%}", "${jndi:x}", "}}{{", "{", "}", "{{ x", "{{Fname|fa}}"]


@pytest.mark.parametrize("hostile", HOSTILE)
def test_m5_braces_in_data_are_literal_everywhere(hostile):
    r = DocxLetterRenderer(TEMPLATES_DIR)
    f = dict(SAMPLES["institute_letter"], Course_Name=f"C {hostile}", To=hostile)
    f["participants"] = [dict(name=hostile, national_id="1", phone="2", email="e@example.test"),
                         dict(name="Second", national_id="3", phone="4", email=f"a{hostile}@example.test")]
    out = r.render("institute_letter", f)
    t = docx_text(out.docx_bytes)
    assert t.count(hostile) >= 3                                  # inserted literally, never interpreted
    assert "۰۰۰۱" in t and t.count("e@example.test") == 1     # nothing substituted from other fields
    assert "" not in t and "" not in t                # shield characters never leak
    assert "[[" not in t


def test_m5_one_hostile_row_does_not_break_a_batch_of_kinds():
    r = DocxLetterRenderer(TEMPLATES_DIR)
    for kind in SAMPLES:
        f = {k: (v if not isinstance(v, str) or k.endswith("gender") else v + " {{Summary}} {{link}}")
             for k, v in SAMPLES[kind].items()}
        out = r.render(kind, f)
        assert out.docx_bytes[:2] == b"PK" and "\ue000" not in docx_text(out.docx_bytes)
        assert "{{Summary}} {{link}}" in docx_text(out.docx_bytes)             # literal text survives in every kind
    hostile_html = r.render("level3_request", dict(SAMPLES["level3_request"], Summary="{{link}}<b>x</b>"))
    assert "{{link}}&lt;b&gt;x&lt;/b&gt;" in hostile_html.html_body


def test_m5_control_characters_are_sanitised():
    r = DocxLetterRenderer(TEMPLATES_DIR)
    f = dict(SAMPLES["level3_request"], Summary="A\x00B\x01C\x0bD\x7fE\x85F\ud800G￾H")
    out = r.render("level3_request", f)
    xml = zipfile.ZipFile(io.BytesIO(out.docx_bytes)).read("word/document.xml")
    assert b"\x00" not in xml and b"\x01" not in xml
    assert "ABCDEFGH" in docx_text(out.docx_bytes)
    Document(io.BytesIO(out.docx_bytes))                          # still a loadable package


# ============================================================= Level-3 rule
def _fact(days, category="Soft Skill"):
    return EnrollmentFact(1, 1, "LMS-1", "C", category, "Evaluation Level 3", date(2026, 9, 30).replace(year=2026) -
                          __import__("datetime").timedelta(days=days), None, None, "Yes", True)


def test_level3_rule_is_ge_60_and_substring_category():
    today = date(2026, 9, 30)
    assert compute_due([_fact(60)], today) and not compute_due([_fact(59)], today)   # documented: >= N
    assert compute_due([_fact(61, "Soft Skills Workshop")], today)                    # substring like Core/legacy SEARCH
    assert not compute_due([_fact(61, "Development")], today)


# ============================================================= SSRF / URL hygiene
@pytest.mark.parametrize("url", ["file:///etc/passwd", "ftp://h.test/x", "gopher://h.test", "//h.test/x", "javascript:1",
                                 "https://u:p@h.test/", "https://h.test/?client_secret=x", "https://h.test/\r\nX: y",
                                 "http://127.0.0.1/", "http://localhost/", "http://169.254.169.254/latest/meta-data",
                                 "http://[::1]/", "http://10.1.2.3/", "http://192.168.0.9/", "http://[fd00:ec2::254]/",
                                 "http://[::ffff:127.0.0.1]/", "http://0.0.0.0/"])
def test_urlcheck_rejects(url):
    with pytest.raises(UnsafeUrlError):
        check_url(url)


def test_urlcheck_allow_private_and_public_ok():
    assert check_url("http://10.1.2.3/x", allow_private_hosts=True)
    assert check_url("http://127.0.0.1:8080/x", allow_private_hosts=True)
    assert check_url("https://ifs.example.invalid/x")                     # unresolvable names pass (connect fails later)
    with pytest.raises(UnsafeUrlError):                                    # explicit opt-in never re-enables bad schemes/creds
        check_url("file:///x", allow_private_hosts=True)
    with pytest.raises(UnsafeUrlError):
        check_url("http://u:p@10.0.0.1/", allow_private_hosts=True)


BASIC = {"auth_mode": "basic", "username": "u", "password": "p"}


def test_ifs_config_rejects_private_by_default_and_allows_when_set():
    for bad in ("http://127.0.0.1:9", "http://169.254.169.254", "http://10.0.0.5/ifs"):
        with pytest.raises(IfsConfigError, match="private|allow_private_hosts"):
            IfsEmployeeSource({"base_url": bad, **BASIC})
    IfsEmployeeSource({"base_url": "http://10.0.0.5/ifs", "allow_private_hosts": True, **BASIC}).close()
    for k, v in (("token_url", "http://169.254.169.254/token"), ("login_url", "file:///x")):
        with pytest.raises(IfsConfigError):
            IfsEmployeeSource({"base_url": "https://ifs.example.invalid", "auth_mode": "oauth2_client_credentials",
                               "client_id": "c", "client_secret": "s", "token_url": "https://t.example.invalid",
                               **{k: v}})
    with pytest.raises(IfsConfigError, match="query"):
        IfsEmployeeSource({"base_url": "https://ifs.example.invalid", "auth_mode": "oauth2_client_credentials",
                           "client_id": "c", "client_secret": "s", "token_url": "http://h.example.invalid/o?apikey=K"})
    with pytest.raises(IfsConfigError, match="same host"):
        IfsEmployeeSource({"base_url": "https://ifs.example.invalid", "employees_path": "https://other.example.invalid/x",
                           **BASIC})


def test_ifs_redirect_is_never_followed_and_authorization_not_forwarded():
    seen = []

    def handler(req: httpx.Request) -> httpx.Response:
        seen.append((str(req.url), req.headers.get("authorization")))
        return httpx.Response(302, headers={"Location": "http://evil.example.invalid/steal"})
    src = IfsEmployeeSource({"base_url": "https://ifs.example.invalid", **BASIC}, transport=httpx.MockTransport(handler))
    with pytest.raises(IfsUpstreamError, match="redirect"):
        list(src.fetch_employees())
    assert len(seen) == 1 and "evil" not in seen[0][0]                    # the second hop was never requested
    ok, msg = src.test_connection()
    assert not ok and "redirect" in msg


def test_ifs_next_link_to_private_or_other_host_refused_and_response_cap():
    def other(req):
        return httpx.Response(200, json={"value": [], "@odata.nextLink": "http://169.254.169.254/latest"})
    src = IfsEmployeeSource({"base_url": "https://ifs.example.invalid", **BASIC}, transport=httpx.MockTransport(other))
    with pytest.raises(IfsUpstreamError, match="different host"):
        list(src.fetch_employees())

    def big(req):
        return httpx.Response(200, content=b"x" * 5000)
    small = IfsEmployeeSource({"base_url": "https://ifs.example.invalid", "max_response_bytes": 1000, **BASIC},
                              transport=httpx.MockTransport(big))
    with pytest.raises(IfsUpstreamError, match="max_response_bytes"):
        list(small.fetch_employees())


def test_ifs_private_target_rejected_at_request_time_even_if_config_was_public():
    """A next-link resolving to a private address is refused even though it is the same origin name."""
    src = IfsEmployeeSource({"base_url": "https://ifs.example.invalid", **BASIC},
                            transport=httpx.MockTransport(lambda r: httpx.Response(200, json={"value": []})))
    with pytest.raises(IfsUpstreamError, match="unsafe URL"):
        src._send("GET", "http://127.0.0.1/x")


def test_graph_url_hygiene():
    base = {"tenant_id": "t", "client_id": "c", "client_secret": "s", "sender": "a@example.test"}
    for k, v in (("authority_url", "http://169.254.169.254"), ("graph_url", "file:///x"),
                 ("graph_url", "https://u:p@graph.example.invalid"), ("authority_url", "http://127.0.0.1:1")):
        with pytest.raises(MessageValidationError):
            GraphSender({**base, k: v})
    GraphSender({**base, "authority_url": "http://127.0.0.1:1", "graph_url": "http://127.0.0.1:1",
                 "allow_private_hosts": True}).test_connection()          # explicit opt-in works


def test_graph_redirect_not_followed_and_no_token_forward():
    seen = []

    def h(req):
        seen.append(str(req.url))
        if req.url.path.endswith("/token"):
            return httpx.Response(200, json={"access_token": "tok", "expires_in": 3600})
        return httpx.Response(307, headers={"Location": "http://evil.example.invalid/x"})
    g = GraphSender({"tenant_id": "t", "client_id": "c", "client_secret": "s", "sender": "a@example.test",
                     "authority_url": "https://login.example.invalid", "graph_url": "https://graph.example.invalid"},
                    transport=httpx.MockTransport(h))
    res = g.send(mail(to=["a@example.test"], cc=[]))
    assert not res[0].ok and "307" in res[0].error and not any("evil" in u for u in seen)


# ============================================================= streaming IFS sync
@pytest.fixture()
def srv():
    with MockIfsServer() as s:
        yield s


def _ifs(srv, **kw):
    return IfsEmployeeSource({"base_url": srv.base_url, "auth_mode": "oauth2_client_credentials", "client_id": CLIENT_ID,
                              "client_secret": CLIENT_SECRET, "token_url": "/oauth/token", "page_size": 5,
                              "allow_private_hosts": True, **kw})


def test_fetch_is_lazy_generator_honours_max_records_cancel_progress(srv):
    with _ifs(srv) as src:
        it = src.fetch_employees()
        assert iter(it) is it                                             # generator, not a list
        assert srv.state.log.count("GET /PersonInfo") == 0                # nothing fetched until iterated
        next(it)
        assert srv.state.log.count("GET /PersonInfo") == 1                # only the first page so far
        it.close()

        got = list(src.fetch_employees(max_records=7))
        assert len(got) == 7 and src.truncated and "max_records" in src.last_report.errors[-1]

        calls = {"n": 0}

        def cancel():
            calls["n"] += 1
            return calls["n"] > 3
        got = list(src.fetch_employees(should_cancel=cancel))
        assert len(got) == 3 and src.cancelled and not src.truncated

        prog = []
        assert len(list(src.fetch_employees(on_progress=lambda n, p: prog.append((n, p))))) == 25
        assert prog == [(5, 1), (10, 2), (15, 3), (20, 4), (25, 5), (25, 6)][:len(prog)] and prog[-1][0] == 25

        chunks = list(src.fetch_employee_chunks(chunk_size=10))
        assert [len(c) for c in chunks] == [10, 10, 5]


def test_ifs_overlong_value_truncated_with_row_error_no_abort():
    long = "X" * 5000

    def h(req):
        return httpx.Response(200, json={"value": [
            {"EmpNo": "E1", "FirstName": long, "LastName": "L"}, {"EmpNo": "E2", "FirstName": "F", "LastName": "L"}]})
    src = IfsEmployeeSource({"base_url": "https://ifs.example.invalid", **BASIC}, transport=httpx.MockTransport(h))
    emps = list(src.fetch_employees())
    assert len(emps) == 2 and len(emps[0].first_name_en) == 300
    errs = " ".join(src.last_report.errors)
    assert "first_name_en" in errs and "X" * 20 not in errs


# ============================================================= Excel robustness
def _zip_with(members: dict[str, bytes]) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in members.items():
            z.writestr(k, v)
    return b.getvalue()


def test_zip_bomb_guard(tmp_path):
    lim = ExcelLimits(max_uncompressed_bytes=5 * 1024 * 1024)
    bomb = _zip_with({"xl/worksheets/sheet1.xml": b"0" * (8 * 1024 * 1024)})
    assert len(bomb) < 100_000
    with pytest.raises(ExcelFormatError, match="zip bomb"):
        check_zip_bomb(bomb, lim)
    p = tmp_path / "bomb.xlsx"
    p.write_bytes(bomb)
    with pytest.raises(ExcelFormatError, match="decompressed"):
        OpenpyxlExcelGateway(lim).read_idp(str(p))
    ratio = _zip_with({"a": b"0" * (30 * 1024 * 1024)})
    with pytest.raises(ExcelFormatError, match="ratio"):
        check_zip_bomb(ratio, ExcelLimits(max_ratio=50))
    with pytest.raises(ExcelFormatError, match="too many"):
        check_zip_bomb(_zip_with({f"f{i}": b"x" for i in range(30)}), ExcelLimits(max_members=10))
    with pytest.raises(ExcelFormatError, match="unsafe"):
        check_zip_bomb(_zip_with({"../evil": b"x"}), ExcelLimits())
    with pytest.raises(ExcelFormatError, match="too large"):
        p.write_bytes(b"0" * 2000)
        OpenpyxlExcelGateway(ExcelLimits(max_file_bytes=1000)).read_idp(str(p))


def test_row_cap_and_column_cap_and_long_cells(tmp_path):
    rows = [idp_row(i) for i in range(1, 51)]
    rows[3][4] = "S" * 5000
    p = make_idp(tmp_path / "a.xlsx", rows)
    out, rep = OpenpyxlExcelGateway(ExcelLimits(max_rows=20, max_cell_chars=100)).read_idp_report(str(p))
    assert len(out) < 20 and any("row limit" in e for e in rep.errors)
    out, rep = OpenpyxlExcelGateway(ExcelLimits(max_cell_chars=100)).read_idp_report(str(p))
    assert len(out) == 50 and len(out[3].summary) == 100
    assert any("Summary" in e and "truncated" in e for e in rep.errors) and not any("SSSS" in e for e in rep.errors)


def test_cancel_and_time_budget(tmp_path):
    p = make_idp(tmp_path / "a.xlsx", [idp_row(i) for i in range(1, 1300)])
    with pytest.raises(ExcelReadAborted):
        OpenpyxlExcelGateway(should_cancel=lambda: True).read_idp(str(p))
    with pytest.raises(ExcelReadAborted, match="time budget"):
        OpenpyxlExcelGateway(ExcelLimits(max_seconds=-1)).read_idp(str(p))
    assert len(OpenpyxlExcelGateway().read_idp(str(p))) == 1299


def test_formulas_read_as_values_and_xlsm_macros_never_loaded(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "IDP"
    from .synth import IDP_HEADER
    ws.append(IDP_HEADER)
    ws.append(idp_row(1))
    ws.append(["=1+1", "=HYPERLINK(\"http://evil.example\",\"x\")", "5", "CR", "=A2", "-", "Submit Motivation", "Within Plan",
               "Soft Skill", "-", "-", "-", "LMS-100002", "-"])
    p = tmp_path / "f.xlsx"
    wb.save(p)
    out = OpenpyxlExcelGateway().read_idp(str(p))
    assert len(out) == 1 or all(not r.emp_no.startswith("=") for r in out)   # formulas are never returned as text
    assert all(not r.summary.startswith("=") and not r.reporter_name.startswith("=") for r in out)
    # .xlsm with a macro part: read as a plain package, vbaProject.bin ignored (keep_vba is off)
    data = p.read_bytes()
    z = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            zo.writestr(it, zi.read(it.filename))
        zo.writestr("xl/vbaProject.bin", b"\xd0\xcf\x11\xe0 not really a macro")
    m = tmp_path / "m.xlsm"
    m.write_bytes(z.getvalue())
    assert len(OpenpyxlExcelGateway().read_idp(str(m))) == len(out)


def test_one_bad_row_never_aborts_the_read(tmp_path, monkeypatch):
    from app.integrations.excel import normalize as N
    real = N.clean_ticket

    def boom(v):
        if "100003" in str(v):
            raise RuntimeError("secret-cell-value")
        return real(v)
    monkeypatch.setattr(N, "clean_ticket", boom)
    p = make_idp(tmp_path / "a.xlsx", [idp_row(i) for i in range(1, 6)])
    out, rep = OpenpyxlExcelGateway().read_idp_report(str(p))
    assert len(out) == 4 and rep.skipped == 1
    assert rep.errors == ["IDP row 8: could not be parsed (RuntimeError)"] or "could not be parsed" in rep.errors[0]
    assert "secret-cell-value" not in " ".join(rep.errors)


def test_write_legacy_neutralises_formula_injection(tmp_path):
    from app.contracts import LegacyData, RunRecord
    evil = ['=HYPERLINK("http://evil.example","x")', "+1+1", "-2+3", "@SUM(A1)"]
    runs = [RunRecord(e, e, "in_person", e, 8.0, date(2027, 1, 1), date(2027, 1, 2), e, e, 1, None, None, None, "-", "")
            for e in evil]
    p = tmp_path / "out.xlsx"
    OpenpyxlExcelGateway().write_legacy(LegacyData(runs=runs), str(p))
    wb = load_workbook(p)                                            # formulas would show data_type 'f'
    assert not [c.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]
    back = OpenpyxlExcelGateway().read_legacy(str(p))
    assert [r.course_name for r in back.runs] == evil                 # value preserved as text, idempotent


# ============================================================= e-mail: per recipient + idempotency
@pytest.fixture()
def smtpd():
    h = Handler()
    c = Controller(h, hostname="127.0.0.1", port=free_port(), authenticator=authenticator, auth_required=False,
                   auth_require_tls=False)
    c.start()
    yield h, c
    c.stop()


COHORT = [f"person{i}@example.test" for i in range(1, 6)]


def test_each_recipient_gets_own_message_with_only_own_address(smtpd):
    import email as stdemail
    import email.policy
    h, c = smtpd
    res = SmtpSender(smtp_cfg(c)).send(mail(to=COHORT, cc=[], attachments=[]))
    assert [r.recipient for r in res] == COHORT and all(r.ok for r in res)
    assert len({r.message_id for r in res}) == 5                       # own Message-ID each
    assert len(h.messages) == 5
    for (_frm, rcpts, content), who in zip(h.messages, COHORT, strict=True):
        raw = content.decode("utf-8", "replace")
        assert rcpts == [who]
        for other in COHORT:
            assert (other in raw) == (other == who), f"{who}'s message leaks {other}"
        m = stdemail.message_from_bytes(content, policy=stdemail.policy.default)
        assert str(m["To"]) == who and m["Cc"] is None and m["Bcc"] is None


def test_refusal_of_one_recipient_does_not_affect_others_and_cc_is_explicit(smtpd):
    h, c = smtpd
    to = ["ok1@example.test", "refuse-me@example.test", "ok2@example.test"]
    res = SmtpSender(smtp_cfg(c)).send(mail(to=to, cc=["boss@example.test"], attachments=[]))
    by = {r.recipient: r for r in res}
    assert by["ok1@example.test"].ok and by["ok2@example.test"].ok and not by["refuse-me@example.test"].ok
    assert "550" in by["refuse-me@example.test"].error
    delivered = [r for _, r, _ in h.messages]
    assert len(delivered) == 3 and all("boss@example.test" in d for d in delivered)     # cc applied per message, only because requested
    none_cc = SmtpSender(smtp_cfg(c)).send(mail(to=["x@example.test"], cc=[], attachments=[]))
    assert none_cc[0].ok and h.messages[-1][1] == ["x@example.test"]


def test_dry_run_per_recipient_and_idempotency_per_recipient(smtpd):
    h, c = smtpd
    s = SmtpSender(smtp_cfg(c))
    res = s.send(mail(to=COHORT, cc=[], attachments=[]), dry_run=True)
    assert [r.recipient for r in res] == COHORT and all(r.message_id == "dry-run" for r in res) and not h.messages
    m = mail(to=COHORT[:3], cc=[], attachments=[], idempotency_key="cohort-key-1")
    s.send(m)
    assert len(h.messages) == 3
    bigger = mail(to=COHORT[:3], cc=[], attachments=[], idempotency_key="cohort-key-1")
    s.send(bigger)
    assert len(h.messages) == 3                                        # replay: nobody gets a second copy


def test_idempotency_conflict_on_any_content_change_and_key_is_hashed():
    sent = []

    class Fake(SmtpSender):
        def _deliver(self, to, cc, subject, msg):
            sent.append(to[0])
            from app.contracts import SendResult
            return [SendResult(a, True, "id") for a in list(to) + list(cc)]
    s = Fake({"host": "h", "from_addr": "a@example.test"})
    base = dict(to=["a@example.test"], subject="S", html_body="<p>b</p>", cc=(), attachments=(("f.txt", b"1", "text/plain"),),
                idempotency_key="key-abcdef-1")
    assert s.send(OutgoingEmail(**base))[0].ok
    assert s.send(OutgoingEmail(**base))[0].ok and sent == ["a@example.test"]
    for change in ({"subject": "S2"}, {"cc": ("c@example.test",)}, {"html_body": "<p>other</p>"},
                   {"attachments": (("f.txt", b"2", "text/plain"),)}, {"attachments": (("g.txt", b"1", "text/plain"),)},
                   {"to": ["a@example.test", "b@example.test"]}):
        with pytest.raises(IdempotencyConflict):
            s.send(OutgoingEmail(**{**base, **change}))
    assert sent == ["a@example.test"]                                  # nothing extra was sent
    store = s.store
    assert hash_key("key-abcdef-1") in store._d and "key-abcdef-1" not in store._d      # raw key never stored


def test_pending_claim_blocks_then_expires_and_can_be_retried():
    now = [1000.0]
    calls = []

    class Flaky(SmtpSender):
        def _deliver(self, to, cc, subject, msg):
            calls.append(1)
            from app.contracts import SendResult
            return [SendResult(a, True, "id") for a in list(to) + list(cc)]
    store = InMemoryIdempotencyStore()
    s = Flaky({"host": "h", "from_addr": "a@example.test"}, idempotency_store=store, wall_clock=lambda: now[0],
              pending_ttl_s=60)
    msg = OutgoingEmail(to=["a@example.test"], subject="S", html_body="b", idempotency_key="pending-key-1")
    from app.integrations.email.common import fingerprint
    store.claim(hash_key("pending-key-1"), fingerprint(msg), now[0], 60)      # a crashed sender left this behind
    with pytest.raises(IdempotencyInProgress):
        s.send(msg)
    assert not calls
    now[0] += 61                                                              # TTL passed: stale, retry allowed
    assert s.send(msg)[0].ok and len(calls) == 1
    assert s.send(msg)[0].ok and len(calls) == 1                              # and now it is a replay
    assert isinstance(IdempotencyInProgress("x"), IdempotencyConflict)


def test_failed_send_releases_the_key_for_immediate_retry():
    class Down(SmtpSender):
        ok = False

        def _deliver(self, to, cc, subject, msg):
            from app.contracts import SendResult
            if not self.ok:
                raise ConnectionError("down")
            return [SendResult(a, True, "id") for a in list(to) + list(cc)]
    s = Down({"host": "h", "from_addr": "a@example.test"})
    msg = OutgoingEmail(to=["a@example.test"], subject="S", html_body="b", idempotency_key="retry-key-01")
    assert not s.send(msg)[0].ok
    s.ok = True
    assert s.send(msg)[0].ok                                                  # not stuck in 'pending'


# ============================================================= misc integrations findings
def test_l5_time_range_and_mixed_script_days_are_bidi_isolated_in_english_section():
    r = DocxLetterRenderer(TEMPLATES_DIR)
    f = dict(SAMPLES["course_announcement"], Day="یکشنبه و دوشنبه")
    en = docx_text(r.render("course_announcement", f, "en").docx_bytes)
    assert "⁨یکشنبه و دوشنبه⁩ | ‪10:00-16:00‬" in en
    html = r.render("course_announcement", f, "en").html_body
    assert "⁨" in html and "‪10:00-16:00‬" in html


def test_watcher_ignores_partial_writes_and_reports_stable_file(tmp_path):
    from app.integrations.excel import FolderWatcher
    t = [0.0]
    w = FolderWatcher(tmp_path, debounce_s=2, clock=lambda: t[0])
    f = tmp_path / "x.xlsx"
    f.write_bytes(b"PK\x03\x04partial")
    for t[0] in (0, 3, 6):
        assert w.poll() == []
    make_ifs(f, [ifs_row(1)])
    for t[0] in (10, 13):
        r = w.poll()
    assert r == [f]


def test_factory_excel_source_uses_limits_and_reports_no_pii(tmp_path):
    p = make_ifs(tmp_path / "e.xlsx", [ifs_row(1, **{"Emp No": ""}), ifs_row(2)])
    src = factory.build_employee_source("excel", {"path": str(p)})
    assert len(list(src.fetch_employees())) == 1 and datetime.now(timezone.utc)
    assert not any("EMP" in e for e in src.last_report.errors)
