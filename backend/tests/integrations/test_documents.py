import io
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.documents import SPECS, DocxLetterRenderer, MissingFieldError, TemplateError, engine
from app.documents.html import build_html

from .samples import SAMPLES, TEMPLATES_DIR

R = DocxLetterRenderer(TEMPLATES_DIR)
KINDS = list(SPECS)


def text_of(docx_bytes: bytes) -> str:
    d = Document(io.BytesIO(docx_bytes))
    out = []
    for root in engine.story_roots(d):
        out += [engine.paragraph_text(p) for p in engine.all_paragraphs(root)]
    return "\n".join(out)


@pytest.mark.parametrize("kind", KINDS)
def test_render_every_kind(kind):
    r = R.render(kind, SAMPLES[kind])
    assert r.kind == kind and r.filename.endswith(".docx") and r.subject and "\n" not in r.subject
    t = text_of(r.docx_bytes)
    assert "{{" not in t and "}}" not in t and "[[LANG" not in t
    assert r.docx_bytes[:2] == b"PK" and "<table" in r.html_body and "<script" not in r.html_body.lower()
    assert " class=" not in r.html_body                                  # inline styles only
    assert "{{" not in r.html_body


def test_required_fields_and_lang():
    for k in KINDS:
        req = list(R.required_fields(k))
        assert req and len(req) == len(set(req))
    fa = R.required_fields("course_announcement", "fa")
    en = R.required_fields("course_announcement", "en")
    assert "Farsi_Course_Name" in fa and "Course_Name" not in fa
    assert "Course_Name" in en and "Farsi_Course_Name" not in en
    assert "participants" in R.required_fields("institute_letter")


@pytest.mark.parametrize("kind", KINDS)
def test_strict_missing_field_named(kind):
    for name in [f for f in R.required_fields(kind)]:
        fields = dict(SAMPLES[kind])
        fields.pop(name, None)
        # supply alternatives the renderer derives from, then drop the target itself
        with pytest.raises(MissingFieldError) as ei:
            R.render(kind, {k: v for k, v in fields.items() if k != name} if name in SAMPLES[kind] else
                     _without_derivation(kind, name))
        assert name in ei.value.fields and name in str(ei.value)


def _without_derivation(kind, name):
    """Sample fields where *name* is absent even though it would normally be derived from a raw field."""
    derived_from = {f.name: f.derived_from for f in SPECS[kind].fields}.get(name, "")
    fields = dict(SAMPLES[kind])
    for raw in re.split(r"[ /()]+", derived_from):
        fields.pop(raw.replace("len", "").strip(), None)
    for alias in ("gender", "manager_gender", "reporter_gender"):
        if name in ("Ftitle", "En_title", "Manager_fTitle", "Manager_Title", "Reporter_F_tilte", "Reporter_title"):
            fields.pop(alias, None)
    if name == "per_attendee":
        fields.pop("invoice_total", None)
        fields.pop("attendees", None)
    if name == "attendee_count":
        fields.pop("attendees", None)
    fields.pop(name, None)
    return fields


@pytest.mark.parametrize("bad", ["", "   ", None])
def test_blank_is_missing_not_silently_empty(bad):
    f = dict(SAMPLES["level3_request"], Summary=bad)
    with pytest.raises(MissingFieldError, match="Summary"):
        R.render("level3_request", f)


def test_list_item_errors_name_the_item():
    f = dict(SAMPLES["institute_letter"])
    f["participants"] = [dict(f["participants"][0]), dict(f["participants"][1], phone="")]
    with pytest.raises(MissingFieldError) as ei:
        R.render("institute_letter", f)
    assert ei.value.fields == ["participants[2].phone"]
    f["participants"] = []
    with pytest.raises(MissingFieldError, match="participants"):
        R.render("institute_letter", f)
    f["participants"] = "oops"
    with pytest.raises(MissingFieldError):
        R.render("institute_letter", f)


def test_rows_repeat_and_content():
    f = dict(SAMPLES["institute_letter"])
    f["participants"] = [dict(name=f"P{i}", national_id=f"00000000{i:02d}", phone="0912", email=f"p{i}@example.test")
                         for i in range(7)]
    r = R.render("institute_letter", f)
    d = Document(io.BytesIO(r.docx_bytes))
    tbl = d.tables[0]
    assert len(tbl.rows) == 8            # header + 7
    assert [c.text for c in tbl.rows[7].cells] == ["P6", "0000000006", "0912", "p6@example.test"]
    assert r.html_body.count("p3@example.test") == 1


def test_settlement_derivations_and_money():
    r = R.render("settlement_letter", SAMPLES["settlement_letter"])
    assert r.fields["attendee_count"] == "5" and r.fields["per_attendee"] == "40,500,000"
    assert r.fields["invoice_total"] == "202,500,000"
    assert r.fields["start_date_gregorian"] == "2024/12/06" and r.fields["start_date_jalali"] == "1403/09/16"
    assert r.fields["attendees"][0]["label"] == "Example Course (EX100)-MTC-000001-Test Person1"
    t = text_of(r.docx_bytes)
    assert "Example Course (EX100)-MTC-000001-Test Person5" in t and "202,500,000" in t
    f = dict(SAMPLES["settlement_letter"], invoice_total="100")       # 3 attendees can't divide evenly
    f["attendees"] = f["attendees"][:3]
    assert R.render("settlement_letter", f).fields["per_attendee"] == "34"    # first share, sum stays exact


def test_dates_derived_for_announcement():
    r = R.render("course_announcement", SAMPLES["course_announcement"])
    v = r.fields
    assert v["jStart_Date"] == "1402/03/09" and v["Date_Horoof"] == "نهم خرداد یک هزار و چهارصد و دو"
    assert v["start_Date"] == "2023/05/30" and v["Deadline"] == "Friday-26-May-2023"
    assert v["deadline_farsi"] == "5 خرداد 1402" and v["Time_Farsi"] == "12 ساعت" and v["Duration"] == "12 Hours"
    assert v["guide_link"] == v["link"]
    assert r.filename == "course_announcement_both.docx"


def test_explicit_fields_win_over_derived():
    f = dict(SAMPLES["course_announcement"], jStart_Date="1999/01/01", Deadline="Custom")
    v = R.render("course_announcement", f).fields
    assert v["jStart_Date"] == "1999/01/01" and v["Deadline"] == "Custom"


def test_language_selection():
    both = text_of(R.render("course_announcement", SAMPLES["course_announcement"]).docx_bytes)
    fa = R.render("course_announcement", SAMPLES["course_announcement"], "fa")
    en = R.render("course_announcement", SAMPLES["course_announcement"], "en")
    assert "Training Course Registration" in both and "ثبت نام در دوره آموزشی" in both
    tf, te = text_of(fa.docx_bytes), text_of(en.docx_bytes)
    assert "ثبت نام" in tf and "Training Course Registration" not in tf
    assert "Training Course Registration" in te and "ثبت نام" not in te
    assert fa.filename.endswith("_fa.docx") and en.subject.startswith("Training Course")
    assert "dir=\"ltr\"" not in fa.html_body.split("<body")[1].split("<table")[0]
    f = {k: v for k, v in SAMPLES["course_announcement"].items() if k in
         ("Course_Name", "Type", "Day", "Time", "Sessions_Count", "link", "start_date", "deadline_date",
          "duration_hours")}
    assert R.render("course_announcement", f, "en")                       # fa fields not needed for en
    with pytest.raises(MissingFieldError):
        R.render("course_announcement", f, "both")
    with pytest.raises(ValueError):
        R.render("course_announcement", f, "de")  # type: ignore[arg-type]


def test_persian_digits_option():
    f = SAMPLES["course_announcement"]
    on = text_of(R.render("course_announcement", f, "fa").docx_bytes)
    off = text_of(DocxLetterRenderer(TEMPLATES_DIR, persian_digits=False).render("course_announcement", f, "fa").docx_bytes)
    over = text_of(R.render("course_announcement", dict(f, persian_digits=False), "fa").docx_bytes)
    assert "۱۴۰۲/۰۳/۰۹" in on and "1402/03/09" not in on
    assert "1402/03/09" in off and "۱۴۰۲" not in off.split("تاریخ شروع")[1][:40]
    assert over == off
    # English section never gets Persian digits
    assert "2023/05/30" in text_of(R.render("course_announcement", f).docx_bytes)
    # time range is isolated LTR so it is not reversed inside RTL text
    assert "\u202a۱۰:۰۰-۱۶:۰۰\u202c" in on


def test_rtl_and_bidi_preserved():
    r = R.render("institute_letter", SAMPLES["institute_letter"])
    d = Document(io.BytesIO(r.docx_bytes))
    paras = [p for p in d.paragraphs if p.text.strip()]
    assert paras and all(p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None for p in paras)
    assert all(run._r.rPr.find(qn("w:rtl")) is not None for p in paras for run in p.runs if run.text.strip())
    assert d.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    # header/footer letterhead images survive
    assert any("image" in rel.reltype for s in d.sections for rel in s.header.part.rels.values())


def test_hide_if_empty_and_defaults():
    f = dict(SAMPLES["line_manager_approval"])
    f.pop("signer_phone")
    f.pop("signer_email")
    f.pop("signer_title")
    f.pop("signer_name")
    r = R.render("line_manager_approval", f)
    t = text_of(r.docx_bytes)
    assert "Talent Development Team" in t and "Human Resources" in t
    lines = [ln for ln in t.splitlines() if ln.strip()]
    assert lines[-1] == "Human Resources"          # no empty phone/email lines left behind
    assert "+98" not in r.html_body and "<br><br>" not in r.html_body
    g = R.render("settlement_letter", {k: v for k, v in SAMPLES["settlement_letter"].items()})
    assert "Talent Development Team" in text_of(g.docx_bytes)
    i = R.render("institute_letter", {k: v for k, v in SAMPLES["institute_letter"].items() if k != "Attachment"})
    assert "ندارد" in text_of(i.docx_bytes)


def test_line_manager_limit_legacy_fields_and_titles():
    f = dict(SAMPLES["line_manager_approval"])
    f["employees"] = [dict(name=f"E{i}", summary="C", ticket="https://t.example.test/x") for i in range(10)]
    with pytest.raises(ValueError, match="at most 9"):
        R.render("line_manager_approval", f)
    legacy = {k: v for k, v in f.items() if k != "employees"}
    for i in range(1, 10):
        legacy.update({f"Empname{i}": f"E{i}", f"EmpSummary{i}": "C", f"EmpTicket{i}": f"https://t.example.test/{i}"})
    r = R.render("line_manager_approval", legacy)
    assert len(r.fields["employees"]) == 9
    assert len(Document(io.BytesIO(r.docx_bytes)).tables[0].rows) == 10
    r2 = R.render("level3_request", SAMPLES["level3_request"] | {"gender": "female"})
    assert r2.fields["Ftitle"] == "سرکار خانم" and r2.fields["En_title"] == "Ms."
    m = SAMPLES["level3_manager_check"]
    r3 = R.render("level3_manager_check", {k: v for k, v in m.items() if k != "reporter_gender"} |
                  {"Reporter_F_title": "جناب آقای", "Reporter_title": "Mr."})
    assert r3.fields["Reporter_F_tilte"] == "جناب آقای"


def test_html_escaping_and_link_safety():
    f = dict(SAMPLES["level3_request"], Summary='<script>alert(1)</script> & "q"',
             link="javascript:alert(1)")
    r = R.render("level3_request", f)
    assert "<script>" not in r.html_body and "&lt;script&gt;" in r.html_body
    assert "href=\"javascript" not in r.html_body
    assert "<script>" in text_of(r.docx_bytes)                    # docx keeps it as plain text
    ok = R.render("level3_request", SAMPLES["level3_request"])
    assert 'href="https://icare.example.test/browse/LMS-000001"' in ok.html_body
    assert 'dir="rtl"' in ok.html_body and 'dir="ltr"' in ok.html_body


def test_subject_is_single_line_and_filename_safe():
    f = dict(SAMPLES["institute_letter"], Course_Name="Evil\r\nBcc: x@example.test", Letter_No="../../etc/1403/5")
    r = R.render("institute_letter", f)
    assert "\n" not in r.subject and "\r" not in r.subject
    assert "/" not in r.filename and ".." not in r.filename and r.filename.startswith("institute_letter_")


def test_xml_illegal_chars_removed():
    f = dict(SAMPLES["level3_request"], Summary="A\x00B\x0bC")
    assert "ABC" in text_of(R.render("level3_request", f).docx_bytes)


def test_errors_unknown_kind_missing_template_bad_template(tmp_path):
    with pytest.raises(ValueError):
        R.render("nope", {})  # type: ignore[arg-type]
    with pytest.raises(FileNotFoundError):
        DocxLetterRenderer(str(tmp_path)).render("level3_request", SAMPLES["level3_request"])
    d = Document()
    d.add_paragraph("Hello {{undeclared_field}} and {{Summary}} {{link}} {{Fname}} {{Ftitle}} {{En_title}} "
                    "{{Reporter__Name}}")
    d.save(tmp_path / "level3_request.docx")
    with pytest.raises(TemplateError, match="undeclared_field"):
        DocxLetterRenderer(str(tmp_path)).render("level3_request", SAMPLES["level3_request"])


def test_engine_placeholder_split_across_runs_keeps_format():
    d = Document()
    p = d.add_paragraph()
    for chunk, bold in (("Dear ", False), ("{{Na", True), ("me}", False), ("} !", False)):
        r = p.add_run(chunk)
        r.bold = bold
        r._r.get_or_add_rPr().append(__import__("docx").oxml.OxmlElement("w:rtl"))
    engine.replace_everywhere(d, lambda n, f: "Zed" if n == "Name" else None)
    assert p.text == "Dear Zed !"
    assert p.runs[1].bold is True and p.runs[1]._r.rPr.find(qn("w:rtl")) is not None
    assert p.runs[1].text == "Zed" and p.runs[2].text == ""


def test_engine_multiple_placeholders_one_run_and_newlines():
    d = Document()
    p = d.add_paragraph("{{a}}-{{ b }}-{{a|ltr}}")
    engine.replace_everywhere(d, lambda n, f: {"a": "1", "b": "x\ny"}[n])
    assert p.text == "1-x\ny-1"
    assert engine.leftover_placeholders(d) == []
    d2 = Document()
    d2.add_paragraph("broken {{x")
    assert engine.leftover_placeholders(d2) == ["<malformed>"]


def test_apply_filter():
    assert engine.apply_filter("10:00-16:00", "fa", False) == "\u202a10:00-16:00\u202c"
    assert engine.apply_filter("abc 12", "fa", True) == "abc ۱۲"
    assert engine.apply_filter("", "fa", True) == "" and engine.apply_filter("x", None, True) == "x"
    assert engine.apply_filter("+98-21", "ltr", True) == "\u202a+98-21\u202c"


def test_html_builder_bad_block():
    with pytest.raises(ValueError):
        build_html([("en", [("bogus",)])], {}, {}, False, "t")


# ------------------------------------------------------------------ template hygiene (no PII, no macros)
TEMPLATE_FILES = sorted(Path(TEMPLATES_DIR).glob("*.docx"))


def test_six_templates_exist():
    assert {p.name for p in TEMPLATE_FILES} == {s.template for s in SPECS.values()}


@pytest.mark.parametrize("path", TEMPLATE_FILES, ids=lambda p: p.name)
def test_template_has_no_macros_links_or_personal_data(path):
    z = zipfile.ZipFile(path)
    names = " ".join(z.namelist()).lower()
    assert "vbaproject" not in names and "recipientdata" not in names and "vbadata" not in names
    blob = b"".join(z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels"))).decode("utf8", "ignore")
    assert "attachedTemplate" not in blob and "mailMerge" not in blob and "dataSource" not in blob
    assert not re.search(r"[A-Z]:\\\\", blob) and "E:\\" not in blob
    d = Document(path)
    text = "\n".join(engine.paragraph_text(p) for root in engine.story_roots(d) for p in engine.all_paragraphs(root))
    assert not re.search(r"@[A-Za-z0-9-]+\.[a-z]{2,}", text), "e-mail address in template"
    assert not re.search(r"\d{7,}", text.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))), "long number in template"
    assert "MERGEFIELD" not in blob
    spec = next(s for s in SPECS.values() if s.template == path.name)
    used = engine.template_placeholders(d)
    declared = {f.name for f in spec.fields}
    for ls in spec.lists:
        declared |= {f"{ls.name}.{f.name}" for f in ls.item_fields} | {f"{ls.name}._n"}
    assert used <= declared | {"guide_link"}
    required_in_template = {f.name for f in spec.fields if f.required}
    assert required_in_template - used == set() or path.name == "settlement_letter.docx", \
        required_in_template - used


def test_wording_spelling_fixed():
    from app.documents.wording import DOCS
    blob = str(DOCS)
    assert "گفتگو" not in blob and "گغتگو" not in blob and "همکار ان" not in blob
    assert "had been attended" not in blob


def test_letter_defaults_file(tmp_path):
    import json
    import shutil
    d = tmp_path / "t"
    shutil.copytree(TEMPLATES_DIR, d, ignore=shutil.ignore_patterns("assets"))
    shutil.copytree(Path(TEMPLATES_DIR) / "assets", d / "assets")
    base = {k: v for k, v in SAMPLES["institute_letter"].items() if not k.startswith("company_")}
    with pytest.raises(MissingFieldError, match="company_economic_code"):
        DocxLetterRenderer(str(d)).render("institute_letter", base)
    ex = json.loads((Path(TEMPLATES_DIR) / "letter_defaults.example.json").read_text())
    (d / "letter_defaults.json").write_text(json.dumps(ex), encoding="utf-8")
    r2 = DocxLetterRenderer(str(d))
    assert "company_economic_code" not in r2.required_fields("institute_letter")
    out = r2.render("institute_letter", base)
    assert "Sample Street" in text_of(out.docx_bytes)
    over = r2.render("institute_letter", dict(base, company_address="Override Address"))
    assert "Override Address" in text_of(over.docx_bytes)
    (d / "letter_defaults.json").write_text('{"nope": 1}')
    with pytest.raises(TemplateError, match="unknown"):
        DocxLetterRenderer(str(d))
    (d / "letter_defaults.json").write_text("[1]")
    with pytest.raises(TemplateError):
        DocxLetterRenderer(str(d))


def test_html_is_responsive_and_outlook_safe():
    from .samples import SAMPLES as S
    h = R.render("course_announcement", S["course_announcement"]).html_body
    assert "max-width:640px" in h and "<!--[if mso]>" in h
    assert h.count('width="640"') == 1 and h.index('width="640"') < h.index("max-width:640px")   # only inside the MSO ghost table
    assert "position:" not in h and "flex" not in h and "<img" not in h


def test_template_metadata_is_neutral():
    for p in TEMPLATE_FILES:
        cp = Document(p).core_properties
        assert cp.author == "TMS" and cp.last_modified_by == "TMS" and cp.category == "" and cp.keywords == ""
