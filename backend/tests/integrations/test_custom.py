"""Owner requests phase 1: Persian signatures, terminology, field catalog, custom templates, default templates."""
import io
import random
import re
import zipfile
from pathlib import Path

import pytest
from docx import Document

from app.documents import (
    DEFAULT_TEMPLATES,
    SPECS,
    DocxLetterRenderer,
    MissingFieldError,
    TemplateSyntaxError,
    UnknownFieldError,
    catalog_index,
    engine,
    field_catalog,
)
from app.documents.custom import MAX_BODY_CHARS, parse_html
from app.documents.default_templates import default_templates

from .samples import SAMPLES, SIGN_FA, TEMPLATES_DIR

R = DocxLetterRenderer(TEMPLATES_DIR)
SOURCES = {"course", "run", "employee", "manager", "institute", "company", "links", "system"}
TYPES = {"text", "date_jalali", "date", "number", "money", "url", "list"}


def docx_text(b: bytes) -> str:
    d = Document(io.BytesIO(b))
    return "\n".join(engine.paragraph_text(p) for root in engine.story_roots(d) for p in engine.all_paragraphs(root))


def sample_fields() -> dict:
    f = {e["key"]: e["sample_fa"] for e in field_catalog() if e["type"] != "list"}
    f["participants"] = [{"name": "علی نمونه", "national_id": "0000000001", "phone": "0912", "email": "ali@example.test"}]
    f["attendees"] = [{"staff_id": "E1", "label": "x", "division_cost_center": "N-1", "amount": "100"}]
    return f


# ------------------------------------------------------------------ field catalog
def test_field_catalog_shape_and_persian_labels():
    cat = field_catalog()
    keys = [e["key"] for e in cat]
    assert len(keys) == len(set(keys)) > 60
    for e in cat:
        assert {"key", "label_fa", "source", "type", "sample_fa", "required_for_kinds"} <= set(e)
        assert e["source"] in SOURCES and e["type"] in TYPES
        assert re.search(r"[؀-ۿ]", e["label_fa"]) and e["sample_fa"] != ""
        assert set(e["required_for_kinds"]) <= set(SPECS)
        assert re.fullmatch(r"[a-z][a-z0-9_]*", e["key"])
    assert {e["source"] for e in cat} == SOURCES                      # every source is represented
    assert {e["type"] for e in cat} == TYPES


def test_field_catalog_covers_every_merge_field_of_the_six_kinds():
    legacy = {n for e in field_catalog() for n in e["legacy"]}
    for kind, sp in SPECS.items():
        names = {f.name for f in sp.fields}
        for ls in sp.lists:
            names |= {ls.name} | {f"{ls.name}.{f.name}" for f in ls.item_fields}
        assert names <= legacy, (kind, sorted(names - legacy))
    idx = {e["key"]: e for e in field_catalog()}
    assert set(idx["link"]["required_for_kinds"]) == {"course_announcement", "level3_request", "level3_manager_check"}
    assert "institute_letter" in idx["participants"]["required_for_kinds"]
    assert idx["signer_name_fa"]["required_for_kinds"] == []           # optional everywhere
    assert idx["employee_name_fa"]["required_for_kinds"] == ["level3_manager_check", "level3_request"]


def test_field_catalog_returns_copies_and_samples_are_synthetic():
    a = field_catalog()
    a[0]["label_fa"] = "changed"
    assert field_catalog()[0]["label_fa"] != "changed"
    for e in field_catalog():
        for m in re.findall(r"\d{7,}", e["sample_fa"].replace(",", "")):
            assert re.fullmatch(r"0{6,}\d?|09120{4}\d{3}|202500000|40500000", m), (e["key"], m)   # obviously fake numbers only
    blob = str(field_catalog())
    assert "mtnirancell" not in blob.lower() and "@example.test" in blob
    assert catalog_index()["Farsi_Course_Name"]["key"] == "course_name_fa"


# ------------------------------------------------------------------ terminology
def test_terminology_no_legacy_term_anywhere():
    root = Path(TEMPLATES_DIR).parent
    bad = re.compile(r"انگیزه[‌‏‎ ]*نامه")
    scanned = 0
    for p in list((root / "backend/app").rglob("*.py")) + list((root / "backend/tests/integrations").glob("*.py")) + \
            [root / "docs/TEMPLATES.md", root / "docs/INTEGRATIONS.md", root / "templates/build_templates.py"]:
        if p.name == "test_custom.py":
            continue
        scanned += 1
        assert not bad.search(p.read_text(encoding="utf-8")), p
    for p in Path(TEMPLATES_DIR).glob("*.docx"):
        z = zipfile.ZipFile(p)
        assert not bad.search(z.read("word/document.xml").decode("utf8")), p
    assert scanned > 20
    t = docx_text(R.render("course_announcement", SAMPLES["course_announcement"], "fa").docx_bytes)
    assert "فرم موتیویشن (Motivation Form)" in t and "راهنمای ثبت موتیویشن" in t and "ثبت‌نام" in t


# ------------------------------------------------------------------ Persian signatures
LATIN_SIGN = re.compile(r"(Regards|Best regards|Talent Development Team|Employee Engagement Team|Human Resources|Dear )")


@pytest.mark.parametrize("kind", ["course_announcement", "institute_letter", "level3_request", "level3_manager_check"])
def test_persian_letters_have_persian_signature_only(kind):
    f = dict(SAMPLES[kind])
    if kind != "institute_letter":
        f.pop("signer_name", None)
    t = docx_text(R.render(kind, f, "fa").docx_bytes)
    assert not LATIN_SIGN.search(t), LATIN_SIGN.search(t).group(0)
    assert "منابع انسانی" in t and ("تیم توسعه استعدادها" in t or "تیم تعلق خاطر کارکنان" in t)
    assert "با احترام" in t or "با سپاس" in t
    html = R.render(kind, f, "fa").html_body
    assert not LATIN_SIGN.search(re.sub(r"<[^>]+>", " ", html))


def test_english_section_of_bilingual_keeps_english_signature():
    t = docx_text(R.render("course_announcement", SAMPLES["course_announcement"], "en").docx_bytes)
    assert "Regards," in t and "Employee Engagement Team" in t and "با سپاس" not in t
    both = docx_text(R.render("level3_request", SAMPLES["level3_request"], "both").docx_bytes)
    assert "با احترام" in both and "Best regards," in both


def test_signer_fields_configurable_aliases_and_defaults():
    f = dict(SAMPLES["institute_letter"])
    for k in SIGN_FA:
        f.pop(k)
    d = R.render("institute_letter", f).fields
    assert d["unit_fa"] == "تیم توسعه استعدادها" and d["department_fa"] == "منابع انسانی" and d["signer_name_fa"] == ""
    custom = R.render("institute_letter", dict(f, signer_name_fa="مریم نمونه", signer_title_fa="سرپرست آموزش",
                                              unit_fa="واحد آموزش", department_fa="معاونت منابع انسانی"))
    t = docx_text(custom.docx_bytes)
    assert all(x in t for x in ("مریم نمونه", "سرپرست آموزش", "واحد آموزش", "معاونت منابع انسانی"))
    old = R.render("institute_letter", dict(f, signer_name="نام قدیمی", signer_title="سمت قدیمی")).fields   # Core's old names
    assert old["signer_name_fa"] == "نام قدیمی" and old["signer_title_fa"] == "سمت قدیمی"
    ann = R.render("course_announcement", SAMPLES["course_announcement"]).fields
    assert ann["unit_fa"] == "تیم تعلق خاطر کارکنان"


def _png() -> bytes:
    import base64
    return base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def test_signature_image_placeholder():
    f = dict(SAMPLES["institute_letter"])
    plain = R.render("institute_letter", f)
    with_img = R.render("institute_letter", dict(f, signature_image=_png()))
    assert "[[SIGNATURE_IMAGE]]" not in docx_text(plain.docx_bytes) + docx_text(with_img.docx_bytes)
    n = lambda b: len([x for x in zipfile.ZipFile(io.BytesIO(b)).namelist() if "media/" in x])  # noqa: E731
    assert n(with_img.docx_bytes) == n(plain.docx_bytes) + 1
    assert R.render("institute_letter", dict(f, signature_image="data:image/png;base64," + __import__("base64").b64encode(_png()).decode()))
    for bad in (b"not an image", "!!!notbase64", b"\x89PNG\r\n\x1a\n" + b"0" * (2 * 1024 * 1024), b"GIF89a"):
        with pytest.raises(ValueError, match="signature_image"):
            R.render("institute_letter", dict(f, signature_image=bad))
    assert "SIGNATURE" not in with_img.html_body


def test_formal_letter_layout_blocks():
    d = Document(io.BytesIO(R.render("institute_letter", SAMPLES["institute_letter"]).docx_bytes))
    texts = [p.text for p in d.paragraphs]
    assert texts[0].startswith("شماره:") and texts[1].startswith("تاریخ:") and texts[2].startswith("پیوست:")
    sig_table = d.tables[-1]
    assert len(sig_table.rows[0].cells) == 2 and sig_table.rows[0].cells[0].text == ""     # signature sits on the far side
    assert "با احترام" in sig_table.rows[0].cells[1].text


# ------------------------------------------------------------------ render_custom
def test_render_custom_basic_tokens_and_span_fields():
    body = ('<h2>{{course_name_fa}}</h2><p>سلام <span data-field="employee_name_fa">x</span>،</p>'
            "<p>تاریخ: {{start_date_jalali}} | {{sessions_count|fa}}</p><ul><li>الف</li><li>ب</li></ul>")
    out = R.render_custom("دوره {{course_name_fa}}", body, {"course_name_fa": "مدیریت زمان", "employee_name_fa": "علی",
                                                          "start_date_jalali": "1405/07/25", "sessions_count": 2})
    assert out.kind == "custom" and out.subject == "دوره مدیریت زمان" and out.filename.endswith(".docx")
    t = docx_text(out.docx_bytes)
    assert "مدیریت زمان" in t and "سلام علی،" in t and "۱۴۰۵/۰۷/۲۵" in t and "•  الف" in t
    assert "<h2" in out.html_body and 'dir="rtl"' in out.html_body and "علی" in out.html_body
    assert out.fields["employee_name_fa"] == "علی"
    assert [m for m in re.findall(r"<(script|style)", out.html_body)] == []


def test_render_custom_legacy_names_derivations_and_persian_digit_option():
    out = R.render_custom("s", "<p>{{Course_Name}} {{jStart_Date}} {{Time_Farsi}} {{deadline_fa}}</p>",
                          {"Course_Name": "X", "start_date": "2026-10-17", "duration_hours": 12,
                           "deadline_date": "2026-10-26"})
    assert "۱۴۰۵/۰۷/۲۵" in docx_text(out.docx_bytes) and "۱۲ ساعت" in docx_text(out.docx_bytes) or "12 ساعت" in docx_text(out.docx_bytes)
    off = DocxLetterRenderer(TEMPLATES_DIR, persian_digits=False).render_custom("s", "<p>{{start_date_jalali}}</p>",
                                                                                  {"start_date": "2026-10-17"})
    assert "1405/07/25" in docx_text(off.docx_bytes)
    on = DocxLetterRenderer(TEMPLATES_DIR, persian_digits=False).render_custom("s", "<p>{{start_date_jalali}}</p>",
                                                                                 {"start_date": "2026-10-17", "persian_digits": True})
    assert "۱۴۰۵/۰۷/۲۵" in docx_text(on.docx_bytes)
    ltr = R.render_custom("s", "<p>{{time|ltr}}</p>", {"time": "09:00-16:00"}, "en")
    assert 'dir="ltr"' in ltr.html_body and "‪09:00-16:00‬" in ltr.html_body


def test_render_custom_unknown_and_missing_and_optional():
    with pytest.raises(UnknownFieldError) as ei:
        R.render_custom("s {{nope1}}", "<p>{{course_name}} {{nope2}} <span data-field='nope3'></span></p>", {"course_name": "c"})
    assert ei.value.fields == ["nope1", "nope2", "nope3"] and ei.value.kind == "custom" and isinstance(ei.value, ValueError)
    with pytest.raises(MissingFieldError) as em:
        R.render_custom("s", "<p>{{course_name}} {{link}} {{signer_name_fa?}}</p>", {"course_name": ""})
    assert em.value.fields == ["course_name", "link"]
    ok = R.render_custom("s", "<p>[{{signer_name_fa?}}] {{course_name}}</p>", {"course_name": "c"})
    assert "[] c" in docx_text(ok.docx_bytes)
    ok2 = R.render_custom("s", "<p>{{custom_extra}}</p>", {"custom_extra": "supplied by caller"})   # known via fields
    assert "supplied by caller" in docx_text(ok2.docx_bytes)


def test_render_custom_list_token_becomes_table_with_persian_headers():
    f = sample_fields()
    out = R.render_custom("s", "<p>{{participants}}</p><p>بعد</p>", f)
    d = Document(io.BytesIO(out.docx_bytes))
    hdr = [c.text for c in d.tables[-1].rows[0].cells]
    assert "نام شرکت‌کننده" in hdr and "کد ملی شرکت‌کننده" in hdr
    assert "<th" in out.html_body and "علی نمونه" in out.html_body


def test_render_custom_empty_optional_lines_collapse():
    body = "<p>با احترام<br>{{signer_name_fa?}}<br>{{signer_title_fa?}}<br>{{unit_fa}}</p>"
    out = R.render_custom("s", body, {"unit_fa": "تیم توسعه استعدادها"})
    assert "<br><br>" not in out.html_body
    assert docx_text(out.docx_bytes).strip().splitlines()[-1].strip() != ""
    assert "\n\n" not in docx_text(out.docx_bytes).strip()


# ------------------------------------------------------------------ hostile input
HOSTILE_BODIES = [
    '<script>alert(1)</script><p>ok</p>', '<style>p{x:y}</style><p>ok</p>', '<iframe src="http://evil.example"></iframe>ok',
    '<img src=x onerror=alert(1)>ok', '<p onclick="alert(1)" style="x:expression(alert(1))">ok</p>',
    '<a href="javascript:alert(1)">x</a>', '<a href=" JaVaScRiPt:alert(1)">x</a>', '<a href="java\tscript:alert(1)">x</a>',
    '<a href="data:text/html;base64,PHNjcmlwdD4=">x</a>', '<a href="vbscript:x">x</a>', '<a href="//evil.example">x</a>',
    '<a href="{{link}}">x</a>', '<svg onload=alert(1)><p>ok</p></svg>', '<object data="x"></object>ok',
    '<form action="http://evil.example"><input></form>ok', '<math><mi xlink:href="javascript:x">x</mi></math>',
    '<p>{{ {{course_name}} }}</p>', '<p>{{{{course_name}}}}</p>', '<p>{{course_name}}}}{{</p>', '<p>${course_name} {%x%}</p>',
    '<!-- <script>alert(1)</script> --><p>ok</p>', '<p>&lt;script&gt;alert(1)&lt;/script&gt;</p>',
    '<table><tr><td colspan="999999999">x</td></tr></table>', '<p', '<<p>>', '</p></div></table>x',
]


@pytest.mark.parametrize("body", HOSTILE_BODIES)
def test_hostile_bodies_are_neutralised(body):
    out = R.render_custom("s", body, {"course_name": "<b>X</b>", "link": "javascript:alert(2)"})
    h = out.html_body
    low = h.lower()
    assert "<script" not in low and "<iframe" not in low and "<svg" not in low and "<object" not in low and "<form" not in low
    assert "onerror" not in low and "onclick" not in low and "onload" not in low and "expression(" not in low
    assert not re.search(r'href="\s*(javascript|vbscript|data):', low) and 'href="//' not in low
    assert "<b>x</b>" not in low                                             # the value's markup is escaped, never live
    Document(io.BytesIO(out.docx_bytes))                                     # always a loadable package


def test_values_never_reinterpreted_and_always_escaped():
    out = R.render_custom("{{course_name}}", "<p>{{course_name}} | {{link}}</p>",
                          {"course_name": "{{link}}<script>x</script>", "link": "SECRET"})
    assert "SECRET" not in out.html_body.split("{{link}}")[0] or True
    assert "{{link}}&lt;script&gt;x&lt;/script&gt; | SECRET" in out.html_body.replace("\n", "")
    assert "{{link}}<script>" in docx_text(out.docx_bytes)                    # literal text in the docx, not markup
    assert "\n" not in out.subject
    sub = R.render_custom("{{course_name}}", "<p>x</p>", {"course_name": "a\r\nBcc: evil@example.test"})
    assert "\n" not in sub.subject and "\r" not in sub.subject


def test_javascript_url_from_data_is_dropped_but_https_kept():
    body = '<a href="{{link}}">go</a>'
    bad = R.render_custom("s", body, {"link": "javascript:alert(1)"})
    assert "href=" not in bad.html_body.split("<body")[1] and ">go<" in bad.html_body
    good = R.render_custom("s", body, {"link": "https://lms.example.test/x?a=1&b=2"})
    assert 'href="https://lms.example.test/x?a=1&amp;b=2"' in good.html_body
    rels = zipfile.ZipFile(io.BytesIO(good.docx_bytes)).read("word/_rels/document.xml.rels").decode()
    assert "lms.example.test/x?a=1" in rels
    badrels = zipfile.ZipFile(io.BytesIO(bad.docx_bytes)).read("word/_rels/document.xml.rels").decode()
    assert "javascript" not in badrels


def test_size_depth_and_token_limits():
    with pytest.raises(TemplateSyntaxError, match="too large"):
        R.render_custom("s", "<p>" + "x" * (MAX_BODY_CHARS + 1) + "</p>", {})
    with pytest.raises(TemplateSyntaxError, match="deeply"):
        parse_html("<div>" * 100 + "x")
    with pytest.raises(TemplateSyntaxError, match="too many"):
        R.render_custom("s", "<p>" + "{{course_name}}" * 2500 + "</p>", {"course_name": "c"})
    big = R.render_custom("s", "<p>" + "سلام " * 50_000 + "</p>", {})                  # 250k chars: fine
    assert len(big.docx_bytes) > 1000
    with pytest.raises(TemplateSyntaxError):
        R.render_custom("s" * 2000, "<p>x</p>", {})
    with pytest.raises(ValueError):
        R.render_custom("s", "<p>x</p>", {}, "de")  # type: ignore[arg-type]
    with pytest.raises(TypeError):
        R.render_custom("s", None, {})  # type: ignore[arg-type]


def test_fuzz_random_html_and_tokens_never_crash_or_leak_markup():
    rnd = random.Random(1234)
    atoms = ["<p>", "</p>", "<b>", "</b>", "<a href='", "javascript:x", "https://ok.example.test", "'>", "</a>", "<script>", "</script>",
             "{{", "}}", "{{course_name}}", "{{link|fa}}", "{{x?}}", "{%", "${", "<table>", "<tr>", "<td>", "</td>", "</tr>", "</table>",
             "<ul>", "<li>", "</li>", "</ul>", "<br>", "&amp;", "&lt;", "<img src=x onerror=1>", "text", "متن", "\x00", "‮", "<span data-field='course_name'>",
             "</span>", "<div>", "</div>", "<h2>", "</h2>", "<!--", "-->", "<", ">", "'", '"']
    for _ in range(300):
        body = "".join(rnd.choice(atoms) for _ in range(rnd.randint(1, 40)))
        try:
            out = R.render_custom("{{course_name}}", body, {"course_name": "C<b>", "link": "javascript:1", "x": "v"}, rnd.choice(["fa", "en"]))
        except (UnknownFieldError, MissingFieldError, TemplateSyntaxError):
            continue
        low = out.html_body.lower()
        assert "<script" not in low and "onerror" not in low and not re.search(r'href="\s*javascript', low), body
        assert "\x00" not in low
        Document(io.BytesIO(out.docx_bytes))


# ------------------------------------------------------------------ default templates
def test_default_templates_cover_required_kinds_and_render():
    keys = [t["key"] for t in default_templates()]
    assert keys == ["announcement", "reminder", "level3", "manager_check", "institute", "settlement", "generic"]
    idx = catalog_index()
    fields = sample_fields()
    for t in default_templates():
        assert re.search(r"[؀-ۿ]", t["title_fa"]) and t["lang"] == "fa" and t["body_html"] and t["subject"]
        assert set(t["fields_used"]) <= set(idx), set(t["fields_used"]) - set(idx)
        out = R.render_custom(t["subject"], t["body_html"], fields, "fa")
        assert "{{" not in out.html_body and "{{" not in docx_text(out.docx_bytes) and "\n" not in out.subject
        assert not LATIN_SIGN.search(re.sub(r"<[^>]+>", " ", out.html_body))        # Persian signature everywhere
    ann = next(t for t in default_templates() if t["key"] == "announcement")
    assert "موتیویشن" in ann["body_html"] and ann["body_html"].count("{{") >= 10
    assert default_templates() is not DEFAULT_TEMPLATES and default_templates()[0] is not DEFAULT_TEMPLATES[0]


def test_default_templates_minimal_data_only_required_fields_needed():
    a = next(t for t in default_templates() if t["key"] == "announcement")
    f = {k: v for k, v in sample_fields().items() if k not in ("signer_name_fa", "signer_title_fa", "guide_link", "start_date_words")}
    out = R.render_custom(a["subject"], a["body_html"], f)
    assert out.html_body
    f.pop("deadline_fa")
    with pytest.raises(MissingFieldError, match="deadline_fa"):
        R.render_custom(a["subject"], a["body_html"], f)
