import datetime as dt
import time

import pytest
from openpyxl import Workbook

from app.contracts import EmployeeRecord, IdpRow, LegacyData, RecordRow, RunRecord
from app.integrations.excel import ExcelFormatError, FolderWatcher, OpenpyxlExcelGateway
from app.integrations.excel import normalize as N

from .synth import idp_row, ifs_row, make_idp, make_ifs

gw = OpenpyxlExcelGateway()


def test_idp_basic(tmp_path):
    rows = [idp_row(1), idp_row(2, sd="-", ed="-", sb="-", att="-"),
            idp_row(3, name="يوسف كريمي [ MTNIrancell - NWG ]", dept="")]
    p = make_idp(tmp_path / "idp.xlsx", rows)
    out, rep = gw.read_idp_report(str(p))
    assert len(out) == 3 and rep.errors == []
    a, b, c = out
    assert a.emp_no == "EMP0001" and a.reporter_name == "Test User1" and a.department == "CR"
    assert a.start_date == dt.date(2023, 5, 30) and a.end_date == dt.date(2023, 5, 31)
    assert a.ticket == "LMS-100001" and a.section_b == "MTC-900001"
    assert b.start_date is None and b.section_b == "-" and b.attendance_status == "-"
    assert c.reporter_name == "یوسف کریمی" and c.department == "NWG"   # yeh/kaf normalised, suffix -> department


def test_idp_header_detection_and_variants(tmp_path):
    rows = [idp_row(1, emp="ABC۰۰۱", ticket="lms ۹۰۰۰۰۲", sd="۱۴۰۲/۰۳/۰۹", ed="2023-06-01", sb="0")]
    p = make_idp(tmp_path / "a.xlsx", rows, header_at=12, dup_header=False)
    (r,) = gw.read_idp(str(p))
    assert r.emp_no == "ABC001" and r.ticket == "LMS-900002"
    assert r.start_date == dt.date(2023, 5, 30) and r.section_b == "-"


def test_idp_errors_have_no_pii(tmp_path):
    rows = [idp_row(1, sd="not-a-date"), idp_row(2, emp=""), idp_row(3, ticket=""),
            idp_row(4, emp="", ticket=""), idp_row(1, sd="2023-05-30"),
            [None] * 14]
    p = make_idp(tmp_path / "a.xlsx", rows)
    out, rep = gw.read_idp_report(str(p))
    assert len(out) == 1
    assert rep.skipped == 4
    joined = " ".join(rep.errors)
    assert "Course Start Date" in joined and "row" in joined
    for secret in ("EMP0001", "Test User", "LMS-1000", "not-a-date", "MTC-"):
        assert secret not in joined


def test_missing_columns_raise_without_values(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["foo", "bar"])
    ws.append(["secret-value", "x"])
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(ExcelFormatError) as ei:
        gw.read_idp(str(p))
    assert "secret-value" not in str(ei.value)


def test_bad_file_types(tmp_path):
    with pytest.raises(FileNotFoundError):
        gw.read_idp(str(tmp_path / "none.xlsx"))
    t = tmp_path / "x.csv"
    t.write_text("a")
    with pytest.raises(ExcelFormatError):
        gw.read_idp(str(t))
    z = tmp_path / "broken.xlsx"
    z.write_bytes(b"not a zip")
    with pytest.raises(ExcelFormatError):
        gw.read_idp(str(z))


def test_employees(tmp_path):
    rows = [ifs_row(1), ifs_row(2, Mobile="", Gender="Female"), ifs_row(1), ifs_row(3, **{"Emp No": ""}),
            ifs_row(4, **{"National Code": "۱۲۳۴۵۶۷۸۹", "Mobile": "+98 935 000 0004", "Division": ""})]
    p = make_ifs(tmp_path / "ifs.xlsx", rows)
    out, rep = gw.read_employees_report(str(p))
    assert [e.emp_no for e in out] == ["EMP0001", "EMP0002", "EMP0004"]
    e = out[0]
    assert isinstance(e, EmployeeRecord)
    assert e.gender == "male" and e.email == "user1@example.test"
    assert e.mobile == "09350000001"            # leading zero restored
    assert e.national_code == "0123456790"      # zero-padded to 10 digits
    assert e.first_name_fa == "علی" and e.last_name_fa == "کریمی"
    assert e.supervisor_emp_no == "EMP0001" and e.status == "Active"
    assert out[1].mobile == "" and out[1].gender == "female"
    assert out[2].mobile == "09350000004" and out[2].division == "Network"
    assert len(rep.errors) == 2 and rep.skipped == 2
    assert not any("EMP" in x for x in rep.errors)


def test_employees_from_named_sheet_auto(tmp_path):
    p = make_ifs(tmp_path / "x.xlsx", [ifs_row(1)], sheet="Employees")
    assert len(gw.read_employees(str(p))) == 1


def _legacy_data():
    emps = [EmployeeRecord("E1", "Ali", "Tester", "علی", "تستی", "male", "a@example.test", "09350000001",
                           "0123456789", "NWG", "23700", "NW-A", "Engineer", "E0", "Active")]
    idp = [IdpRow("E1", "Ali Tester", "5", "CR", "Crucial Conversation", "-", "Submit Motivation", "Within Plan",
                  "Soft Skill", "-", dt.date(2023, 5, 30), dt.date(2023, 5, 31), "LMS-100001", "MTC-1")]
    runs = [RunRecord("Crucial Conversation", "گفتگو در میدان مین", "in_person", "Inst", 12.0,
                      dt.date(2023, 5, 30), dt.date(2023, 5, 31), "Sunday & Monday", "10:00-16:00", 2,
                      dt.date(2023, 5, 24), dt.date(2023, 5, 28), 40500000.0, "MTC-1", "http://x.test")]
    recs = [RecordRow("E1", "LMS-100001", "Crucial Conversation", "Inst", dt.date(2023, 5, 30), "Yes",
                      "Done", "Yes", "Started", "note")]
    return LegacyData(emps, idp, runs, recs)


def test_write_and_read_legacy_roundtrip(tmp_path):
    data = _legacy_data()
    p = tmp_path / "legacy.xlsx"
    gw.write_legacy(data, str(p))
    back, rep = gw.read_legacy_report(str(p))
    assert rep.errors == []
    assert back.employees[0].emp_no == "E1" and back.employees[0].first_name_fa == "علی"
    assert back.employees[0].national_code == "0123456789"
    assert back.idp == data.idp
    assert back.runs == data.runs
    assert back.records == data.records
    # idempotent: writing what we read yields the same data
    p2 = tmp_path / "legacy2.xlsx"
    gw.write_legacy(back, str(p2))
    assert gw.read_legacy(str(p2)).runs == data.runs
    # legacy layout: header rows
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True)
    assert wb.sheetnames == ["IFS", "IDP", "Record", "CHistory"]
    ws = wb["IDP"]
    hdr = [c for c in next(ws.iter_rows(min_row=5, max_row=5, values_only=True))]
    assert hdr[0] == "Employee ID" and hdr[12] == "Ticket"
    assert not list(tmp_path.glob("*.tmp"))


def test_write_rejects_xlsm_and_reader_reads_xlsm(tmp_path):
    with pytest.raises(ExcelFormatError):
        gw.write_legacy(LegacyData(), str(tmp_path / "x.xlsm"))
    src = make_idp(tmp_path / "a.xlsx", [idp_row(1)])
    m = tmp_path / "a.xlsm"
    m.write_bytes(src.read_bytes())          # same package, macro extension
    assert len(gw.read_idp(str(m))) == 1


def test_legacy_without_known_sheets(tmp_path):
    wb = Workbook()
    wb.active.title = "Other"
    p = tmp_path / "o.xlsx"
    wb.save(p)
    with pytest.raises(ExcelFormatError):
        gw.read_legacy(str(p))


def test_source_never_modified(tmp_path):
    p = make_idp(tmp_path / "a.xlsx", [idp_row(i) for i in range(50)])
    before = (p.read_bytes(), p.stat().st_mtime_ns)
    gw.read_idp(str(p))
    gw.read_idp(str(p))
    assert (p.read_bytes(), p.stat().st_mtime_ns) == before
    p.unlink()   # not locked: deletable right after read (matters on Windows)


def test_large_files_fast(tmp_path):
    ifs = make_ifs(tmp_path / "ifs.xlsx", [ifs_row(i) for i in range(1, 3301)])
    idp = make_idp(tmp_path / "idp.xlsx", [idp_row(i) for i in range(1, 2701)])
    t = time.perf_counter()
    emps = gw.read_employees(str(ifs))
    rows = gw.read_idp(str(idp))
    dur = time.perf_counter() - t
    assert len(emps) == 3300 and len(rows) == 2700
    assert dur < 5


def test_wrong_dimension_metadata(tmp_path):
    """Some exporters write <dimension ref="A1"/>; reader must still see all rows."""
    import zipfile
    src = make_idp(tmp_path / "a.xlsx", [idp_row(i) for i in range(5)])
    out = tmp_path / "b.xlsx"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                import re
                data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1"', data)
            zo.writestr(it, data)
    assert len(gw.read_idp(str(out))) == 5


# ------------------------------------------------------------------ normalize
def test_normalize_helpers():
    assert N.split_reporter("Name X [ MTNIrancell - CR ]") == ("Name X", "CR")
    assert N.split_reporter("Plain") == ("Plain", "")
    assert N.parse_date("Sunday-28-May-2023") == (dt.date(2023, 5, 28), True)
    assert N.parse_date("-") == (None, True) and N.parse_date(None) == (None, True)
    assert N.parse_date("2023-13-40") == (None, False)
    assert N.parse_date("1402/13/40") == (None, False)
    assert N.parse_date(45061)[0] == dt.date(2023, 5, 15)
    assert N.parse_date(5)[1] is False
    assert N.parse_date(dt.datetime(2020, 1, 2, 3))[0] == dt.date(2020, 1, 2)
    assert N.parse_date(dt.date(2020, 1, 2))[0] == dt.date(2020, 1, 2)
    assert N.clean_section_b("mtc 123") == "MTC-123" and N.clean_section_b("") == "-"
    assert N.clean_section_b("weird") == "weird"
    assert N.clean_ticket("xx") == "xx"
    assert N.fix_mobile("00989350000001") == "09350000001"
    assert N.fix_national_code(None) == "" and N.fix_national_code(12) == "0000000012"
    assert N.clean_email("nope") == ""
    assert N.gender("مرد") == "male" and N.gender("زن") == "female" and N.gender("x") == ""
    assert N.to_float("1,500") == (1500.0, True) and N.to_float("x") == (None, False)
    assert N.to_int("3") == 3 and N.to_int("x", 7) == 7
    assert N.to_text(3.0) == "3" and N.to_text(True) == "True"
    assert N.to_text(dt.date(2020, 1, 2)) == "2020-01-02"
    assert N.is_blank("  ")


# ------------------------------------------------------------------ watcher
class Clock:
    def __init__(self):
        self.t = 0.0

    def __call__(self):
        return self.t


def test_watcher_debounce_and_dedupe(tmp_path):
    clk = Clock()
    w = FolderWatcher(tmp_path, debounce_s=2, clock=clk)
    assert w.poll() == []
    p = make_idp(tmp_path / "idp.xlsx", [idp_row(1)])
    assert w.poll() == []               # first sight
    clk.t = 1
    assert w.poll() == []               # not stable long enough
    clk.t = 3
    assert w.poll() == [p]              # stable -> reported
    clk.t = 10
    assert w.poll() == []               # not reported twice
    p.touch()                           # mtime changes, content identical
    clk.t = 11
    w.poll()
    clk.t = 14
    assert w.poll() == []               # same hash -> no reprocess
    make_idp(p, [idp_row(1), idp_row(2)])   # real change
    clk.t = 15
    w.poll()
    clk.t = 18
    assert w.poll() == [p]
    p.unlink()
    clk.t = 20
    assert w.poll() == [] and not w._state


def test_watcher_ignores_temp_partial_and_other(tmp_path):
    clk = Clock()
    w = FolderWatcher(tmp_path, debounce_s=1, clock=clk)
    (tmp_path / "~$lock.xlsx").write_bytes(b"x")
    (tmp_path / "a.xlsx.tmp").write_bytes(b"x")
    (tmp_path / "notes.txt").write_text("x")
    (tmp_path / "partial.xlsx").write_bytes(b"PK\x03\x04 truncated")   # still being written
    (tmp_path / "empty.xlsx").write_bytes(b"")
    for t in (0, 2, 4):
        clk.t = t
        assert w.poll() == []


def test_watcher_missing_dir_and_run_loop(tmp_path):
    import threading
    assert FolderWatcher(tmp_path / "nope").poll() == []
    make_idp(tmp_path / "idp.xlsx", [idp_row(1)])
    w = FolderWatcher(tmp_path, debounce_s=0.05)
    seen, stop = [], threading.Event()

    def cb(p):
        seen.append(p.name)
        stop.set()
        raise RuntimeError("callback failure must not kill watcher")

    th = threading.Thread(target=w.run, args=(cb, stop, 0.05))
    th.start()
    th.join(5)
    assert seen == ["idp.xlsx"] and not th.is_alive()
