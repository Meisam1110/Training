"""Synthetic workbook builders (no real data)."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook

from app.integrations.excel.gateway import LEGACY_IFS_COLS

IDP_HEADER = ["Employee ID", "Reporter:  Name", "Job Level", "Reporter Department", "Summary", "Delivery Method",
              "Current Status", "IDP Type", "Course Category", "Attendance Status", "Course Start Date",
              "Course End Date", "Ticket", "Section B"]


def idp_row(i: int, **kw):
    base = {
        "emp": f"EMP{i:04d}", "name": f"Test User{i} [ MTNIrancell - CR ]", "lvl": "5", "dept": "CR",
        "summary": "Crucial Conversation", "dm": "-", "status": "Evaluation Level 3", "typ": "Within Plan",
        "cat": "Soft Skill", "att": "Passed", "sd": "2023-05-30 10:00:00", "ed": "2023-05-31 16:00:00",
        "ticket": f"LMS-{100000 + i}", "sb": "MTC-900001",
    }
    base.update(kw)
    return [base["emp"], base["name"], base["lvl"], base["dept"], base["summary"], base["dm"], base["status"],
            base["typ"], base["cat"], base["att"], base["sd"], base["ed"], base["ticket"], base["sb"]]


def make_idp(path: Path, rows: list[list], header_at: int = 4, dup_header: bool = True, suffix=".xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "IDP"
    ws.cell(1, 1, "exported 2023-05-24")
    for _ in range(header_at - 2):
        ws.append([])
    ws.append(IDP_HEADER)
    if dup_header:
        ws.append(IDP_HEADER)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def ifs_row(i: int, **kw):
    d = {c: "" for c in LEGACY_IFS_COLS}
    d.update({"Emp No": f"EMP{i:04d}", "Employee Status": "Active", "Employment Type": "Local", "Org Code": "NW-A",
              "Organization Name": "Network", "Cost Center": 23700, "Fname": f"Ali{i}", "Lname": f"Tester{i}",
              "Gender": "Male" if i % 2 else "Female", "Mobile": 9350000000 + i, "Email": f"User{i}@Example.Test",
              "Farsi First Name": "علي", "Farsi Last Name": "كريمي", "Position Title": "Engineer",
              "Sup Emp No": "EMP0001", "National Code": 123456789 + i, "Division": "NWG",
              "Date Of Birth": dt.datetime(1990, 1, 1)})
    d.update(kw)
    return [d[c] for c in LEGACY_IFS_COLS]


def make_ifs(path: Path, rows: list[list], sheet="IFS"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(LEGACY_IFS_COLS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path
