"""Legacy workbook / IDP import tests with SYNTHETIC files."""
import os, sys, json, io, zipfile, time
os.environ["QA_BASE"] = "http://127.0.0.1:18701"
from lib import Client
import synth
a = Client("http://127.0.0.1:18701"); r = a.login("qaadmin", "Qa-Admin-Pass-9137!x"); assert r.status_code == 200, r.text
def up(path_or_bytes, name="x.xlsx", ep="/import/legacy-workbook", ct=None):
    data = open(path_or_bytes, "rb").read() if isinstance(path_or_bytes, str) else path_or_bytes
    return a.post(ep, files={"file": (name, data, ct or "application/octet-stream")})
synth.build("/tmp/qa/s1.xlsx", formula_names=True)
r = up("/tmp/qa/s1.xlsx"); print("legacy import 1:", r.status_code, json.dumps(r.json(), ensure_ascii=False)[:900])
r = up("/tmp/qa/s1.xlsx"); j = r.json(); print("legacy import 2 (idempotent?):", r.status_code, {k: j.get(k) for k in ("created", "updated", "unchanged", "skipped", "status")})
emps = a.get("/employees", params={"page_size": 200, "q": "9"}).json()
print("employees:", emps["total"])
en = a.get("/enrollments", params={"page_size": 200}).json(); print("enrollments:", en["total"])
tick = {}
for e in en["items"]: tick.setdefault(e["ticket"], []).append(e["emp_no"])
print("dup tickets:", {k: v for k, v in tick.items() if len(v) > 1})
p = a.get("/employees/9100001?reveal=1").json() if False else None
for q in ("۹۱۰۰۰۰۱", "9100001"):
    r = a.get("/employees", params={"q": q}); print("search", q, r.json()["total"])
r = a.get("/employees/9100001"); print("persian-digit emp:", r.status_code, {k: r.json().get(k) for k in ("first_name_en", "last_name_en", "email", "emp_no")} if r.status_code == 200 else r.text[:100])
r = a.get("/employees", params={"q": "Person"}); 
# unknown employee ticket LMS-300001 present?
print("LMS-300001 (unknown emp) imported:", any(e["ticket"] == "LMS-300001" for e in en["items"]))
print("LMS-100001 owners:", tick.get("LMS-100001"))
print("Terminate rows in enrollments (LMS-200001):", [ (e['emp_no'], e['start_date'], e['current_status']) for e in en['items'] if e['ticket']=='LMS-200001'])
long = a.get("/employees/9200000"); print("long name emp:", long.status_code, len(long.json().get("first_name_en", "")) if long.status_code == 200 else long.text[:100])
# export formula injection
r = a.get("/export/legacy.xlsx"); print("export:", r.status_code, len(r.content))
open("/tmp/qa/export1.xlsx", "wb").write(r.content)
from openpyxl import load_workbook
wb = load_workbook("/tmp/qa/export1.xlsx")
for row in wb["IFS"].iter_rows(min_col=1, max_col=8):
    for c in row:
        if c.data_type == "f": print("FORMULA CELL in export IFS:", c.coordinate, repr(c.value)[:80])
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f" and ws.title != "IFS": print("FORMULA CELL", ws.title, c.coordinate, repr(c.value)[:80])
print("export sheets:", wb.sheetnames)
