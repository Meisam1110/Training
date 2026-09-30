"""Each test asserts the EXPECTED (correct) behaviour. A FAILING test = a reproduced finding (see docs/QA_REPORT.md ids)."""
import datetime as dt, io, zipfile
import pytest
from openpyxl import load_workbook

def _run(hr, **kw):
    cid = hr.post("/courses", json={"name": "QA Course", "category": "Soft Skill"}).json()["id"]
    body = {"course_id": cid, "institute": "QA Inst", "duration_h": 8, "start_date": "2027-03-01", "end_date": "2027-03-02", "days": "Sat", "time": "09:00-16:00",
            "sessions": 2, "deadline": "2027-02-20", "amount": 100000000, "section_b": "MTC-1", "link": "https://x.test/c", **kw}
    r = hr.post("/runs", json=body); assert r.status_code == 201, r.text; return r.json()["id"]

def _participants(hr, rid, n=3):
    seen, ens = set(), []
    for e in hr.get("/enrollments", params={"page_size": 100}).json()["items"]:
        if e["emp_no"] not in seen and len(ens) < n: seen.add(e["emp_no"]); ens.append(e)
    r = hr.post("/enrollments:bulk", json={"enrollment_ids": [e["id"] for e in ens], "run_id": rid, "motivation": "Yes"}); assert r.status_code == 200, r.text
    assert r.json()["updated"], r.json()
    return r.json()["updated"]

def test_QA_B1_settlement_letter_renders_with_real_renderer(hr, seeded):
    """Blocker: Core builds start_date_gregorian/invoice_total/attendees; the real renderer requires start_date_greg/total/rows."""
    rid = _run(hr); _participants(hr, rid)
    r = hr.post("/letters", json={"kind": "settlement_letter", "run_id": rid, "extra": {"invoice_no": "INV-1"}})
    assert r.status_code == 201, r.text

def test_QA_H_level3_loop_closes_after_letter(hr, seeded):
    """High: level3_sent is never written -> an enrollment stays 'due' after its letter was created."""
    from app.models import Enrollment
    from sqlalchemy import select
    import app.main
    sf = hr.c.app.state.session_factory
    with sf() as db:
        e = db.scalars(select(Enrollment)).first(); e.current_status = "Evaluation Level 3"; e.category = "Soft Skill"; e.end_date = dt.date.today() - dt.timedelta(days=90); eid = e.id; db.commit()
    due = lambda: [i for i in hr.get("/automation/due").json()["items"] if i["kind"] == "level3_due" and i["enrollment_id"] == eid]
    assert due()
    assert hr.post("/letters", json={"kind": "level3_request", "enrollment_ids": [eid]}).status_code == 201
    assert not due(), "enrollment still due after Level-3 letter was generated (level3_sent never set)"

def test_QA_H_level3_boundary_consistent_between_dashboard_and_due(hr, seeded):
    """Medium: /automation/due uses >= 60 days, dashboard/candidates use > 60 (legacy: today-end>60)."""
    from app.models import Enrollment
    from sqlalchemy import select
    sf = hr.c.app.state.session_factory
    with sf() as db:
        e = db.scalars(select(Enrollment)).first(); e.current_status = "Evaluation Level 3"; e.category = "Soft Skill"; e.end_date = dt.date.today() - dt.timedelta(days=60); eid = e.id; db.commit()
    in_due = any(i["enrollment_id"] == eid for i in hr.get("/automation/due").json()["items"] if i["kind"] == "level3_due")
    in_dash = any(x["enrollment_id"] == eid for x in hr.get("/dashboard").json()["level3_due"])
    assert in_due == in_dash

def test_QA_M_integer_overflow_id_is_not_500(hr, seeded):
    assert hr.get("/runs/99999999999999999999").status_code in (404, 422)

def test_QA_M_idempotency_key_reuse_with_different_content_is_rejected(hr, admin, seeded, tmp_path):
    admin.put("/connectors/smtp", json={"enabled": True, "config": {"host": "127.0.0.1", "port": 1, "starttls": False, "from_addr": "t@example.test"}})
    rid = _run(hr); _participants(hr, rid)
    lid = hr.post("/letters", json={"kind": "course_announcement", "run_id": rid}).json()["id"]
    b = {"letter_id": lid, "recipients": ["a@example.test"], "idempotency_key": "qa-idem-0001", "dry_run": True}
    assert hr.post("/email/send", json=b).status_code == 200
    r = hr.post("/email/send", json={**b, "subject": "A completely different subject", "cc": ["other@example.test"]})
    assert r.status_code == 409, "same key + different subject/cc silently replays the first result"

def test_QA_M_export_neutralises_spreadsheet_formulas(hr, seeded):
    cid = hr.post("/courses", json={"name": '=HYPERLINK("http://evil.example","x")'}).json()["id"]
    hr.post("/runs", json={"course_id": cid, "start_date": "2026-12-01", "end_date": "2026-12-02"})
    r = hr.get("/export/legacy.xlsx"); assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    formulas = [c.coordinate for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert not formulas, f"formula cells in export: {formulas[:3]}"

def test_QA_M_import_does_not_silently_move_ticket_to_another_employee(hr, admin):
    import sys; sys.path.insert(0, "/home/user/training/qa")
    from conftest import workbook_bytes
    def build(synth, p): synth.build(p, n_emp=5, n_idp=5, edge=False)
    data = workbook_bytes(build)
    hr.post("/import/legacy-workbook", files={"file": ("a.xlsx", data)})
    before = {e["ticket"]: e["emp_no"] for e in hr.get("/enrollments", params={"page_size": 50}).json()["items"]}
    from qa_synth_dup import idp_with_swapped_owner
    r = hr.post("/import/idp", files={"file": ("b.xlsx", idp_with_swapped_owner())}).json()
    after = {e["ticket"]: e["emp_no"] for e in hr.get("/enrollments", params={"page_size": 50}).json()["items"]}
    moved = [t for t in before if before[t] != after.get(t)]
    assert not moved or any(e.get("code") for e in r["errors"]), "ticket ownership changed with no error reported"

def test_QA_M_placeholder_in_data_value_does_not_break_or_inject(hr, seeded):
    """Values containing {{x}} are re-scanned as templates."""
    from app.documents.renderer import DocxLetterRenderer
    r = DocxLetterRenderer("/home/user/training/templates")
    f = dict(Letter_No="LN-1", date="1405/07/08", To="Inst", Course_Name="{{To}}", Jstart_Date="1405/01/01", Days="d", Time="t", Duration="8", Method="m",
             company_economic_code="e", company_national_id="n", company_postal_code="p", company_address="a",
             participants=[{"name": "A", "national_id": "1", "phone": "2", "email": "e@example.test"}])
    res = r.render("institute_letter", f, "fa")
    x = zipfile.ZipFile(io.BytesIO(res.docx_bytes)).read("word/document.xml").decode()
    assert "{{To}}" in x or "{{" in x.replace("&#123;", "{"), "value was interpreted as a template"

def test_QA_M_connector_url_with_embedded_credentials_is_rejected_on_save(admin):
    r = admin.put("/connectors/ifs", json={"enabled": False, "config": {"base_url": "http://user:SuperSecret@127.0.0.1:1", "auth_mode": "basic"}})
    assert r.status_code == 422, "credentials in base_url are stored in plaintext JSON and returned by GET /connectors"

def test_QA_M_manual_attendance_edit_survives_idp_reimport(hr, seeded):
    en = hr.get("/enrollments", params={"page_size": 5}).json()["items"][0]
    hr.patch(f"/enrollments/{en['id']}", json={"attendance": "Attended"})
    hr.post("/import/legacy-workbook", files={"file": ("a.xlsx", seeded["data"])})
    after = hr.get("/enrollments", params={"q": en["ticket"]}).json()["items"][0]
    assert after["attendance"] == "Attended", "manual edit reverted by re-import"

def test_QA_L_null_byte_rejected_cleanly(hr):
    assert hr.post("/courses", json={"name": "a\x00b"}).status_code in (201, 422)   # 500 on PostgreSQL only; see pg_specific.py
